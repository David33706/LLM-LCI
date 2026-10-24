#!/usr/bin/env python3
"""
evaluate.py — Scoring harness for the LCA RAG system.

WHAT THIS DOES
--------------
Runs your evaluation question set through the RAG pipeline and scores it, so you
can tell whether a change (new embedder, hybrid search, reranker, ...) actually
improved things instead of guessing.

It measures three things, from shallow to deep — they only go down, and the drop
between them localises what is broken:

  1. RETRIEVAL  — did the right paper come back, and how high up?
                  Metrics: Hit@k (was the correct DOI in the top-k?) and
                  MRR (Mean Reciprocal Rank — rank 1 scores 1.0, rank 2 -> 0.5, ...).
                  Fast: no LLM needed.

  2. CHUNK-HIT  — were the expected answer's KEY FACTS actually present in the
                  retrieved chunks? The tier BETWEEN "found the paper" and
                  "answered correctly": was the answer even VISIBLE to the model
                  (and to a researcher reading the retrieved text)? A drop from
                  Hit@k -> chunk-hit is a retrieval-coverage gap (unindexed tables,
                  or the right passage not ranking in the top-k). Heuristic
                  key-fact match; needs no LLM, so it runs in EVERY mode.

  3. ANSWERS    — did the generated answer match the expected answer?
                  Three checks, increasing cost:
                    (a) auto    : numeric / substring match (instant, deterministic)
                    (b) judge   : a local LLM decides yes/no (handles phrasing)
                    (c) manual  : side-by-side dump you eyeball
                  Only runs with --generate. A drop from chunk-hit -> answer is the
                  LLM's failure (it had the text and still got it wrong).

TWO MODES (run from the repo root)
---------
  retrieval-only (default):   python eval/evaluate.py --eval-set my_questions.csv
                              (reports paper-hit; chunk-hit too when the eval set has expected_answer)
  full (also answers):        python eval/evaluate.py --eval-set ... --generate \
                                  --llm-model qwen2.5:7b-instruct-q8_0
  full + LLM judge on answers:  ... --generate --judge   (legacy; the adopted judge is
                                  eval/judge_answers.py)

RETRIEVER BACKENDS
------------------
  Defaults: --retriever hybrid --fusion weighted --alpha 0.4 --candidate-k 50 --top-k 10.
  candidate_k stays 50 ON PURPOSE: README's Hit@10 92.0% / MRR 0.671 (v3_50), the
  documented known-answer check, were measured at 50. The runtime pinned in
  step3_orchestrate.py uses 100; pass --candidate-k 100 to score that config.
  --retriever hybrid  (default) : dense + BM25 keyword, fused weighted (default) or RRF
                                  (requires ingest/build_bm25_index.py to have been run;
                                   --bm25-dir defaults to bm25_index/<collection>)
  --retriever dense             : Chroma dense vectors only (the June baseline)
  An empty or uncountable vector store is an ERROR (exit 1), not a warning.

OUTPUTS
-------
  - <out>/results_<timestamp>.json : full detail (every retrieved chunk, scores)
  - <out>/scorecard_<timestamp>.csv: one row per question, human-readable, diff-able
  - a summary table printed to the terminal

The eval set may be .xlsx (the workbook we built — it reads the 'Questions' sheet)
or .csv, with columns: id, question, type, expected_answer, source_doi, source_location, notes.
The default, eval/LCA_RAG_eval_set_v3_50.xlsx (the project's 50-question set), is kept
in the author's private archive and not published. Pass --eval-set with your own
questions file (required columns: id, type, question, source_doi); a missing eval set
is an ERROR (exit 1). Without expected_answer, chunk-hit is skipped and --generate
exits with an ERROR.
"""

import argparse
import csv
import json
import os
import re
import sys
import time
from dataclasses import dataclass, field, asdict
from datetime import datetime
from typing import Any, Dict, List, Optional

# same message as step3_orchestrate.EVAL_SET_MISSING (this script does not import step3)
EVAL_SET_MISSING = ("ERROR: eval set not found: {}. The project's 50-question eval set is not published; "
                    "pass --eval-set with your own .xlsx (sheet \"Questions\") or .csv with columns "
                    "id,type,question,source_doi (optional: valid_dois, expected_answer, source_location).")


# ----------------------------------------------------------------------
# Loading the eval set
# ----------------------------------------------------------------------
@dataclass
class EvalQuestion:
    id: str
    question: str
    type: str                 # NUMBER | PROSE | HARD | TASK1_LOCATE | TASK2_ASSUMPTION ...
    expected_answer: str
    source_doi: str
    source_location: str = ""
    notes: str = ""


def load_eval_set(path: str) -> List[EvalQuestion]:
    if not os.path.exists(path):
        sys.exit(EVAL_SET_MISSING.format(path))
    rows: List[Dict[str, Any]] = []
    ext = os.path.splitext(path)[1].lower()

    if ext in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ImportError("openpyxl needed to read .xlsx. pip install openpyxl") from exc
        wb = load_workbook(path, data_only=True)
        ws = wb["Questions"] if "Questions" in wb.sheetnames else wb[wb.sheetnames[0]]
        header = [str(c.value).strip().lower() if c.value is not None else "" for c in ws[1]]
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r is None or all(v is None for v in r):
                continue
            row = {header[i]: (r[i] if i < len(r) else None) for i in range(len(header))}
            rows.append(row)
    elif ext == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            for row in reader:
                rows.append({(k or "").strip().lower(): v for k, v in row.items()})
    else:
        raise ValueError(f"Unsupported eval-set format: {ext} (use .xlsx or .csv)")

    questions: List[EvalQuestion] = []
    for i, row in enumerate(rows, start=1):
        q = str(row.get("question") or "").strip()
        doi = str(row.get("source_doi") or "").strip()
        if not q or not doi:
            continue  # skip blank/incomplete rows
        questions.append(
            EvalQuestion(
                id=str(row.get("id") or i).strip(),
                question=q,
                type=str(row.get("type") or "").strip().upper() or "UNSPEC",
                expected_answer=str(row.get("expected_answer") or "").strip(),
                source_doi=doi,
                source_location=str(row.get("source_location") or "").strip(),
                notes=str(row.get("notes") or "").strip(),
            )
        )
    if not questions:
        raise ValueError(f"No usable questions found in {path}")
    return questions


# ----------------------------------------------------------------------
# DOI matching (retrieval scoring)
# ----------------------------------------------------------------------
def normalize_doi(doi: Optional[str]) -> str:
    """Normalize DOIs so '10.1016/j.afres.2025.100716', with/without https prefix,
    and the underscore filename form '10_1016_j_afres_2025_100716' all compare equal."""
    if not doi:
        return ""
    d = str(doi).strip().lower()
    d = re.sub(r"^https?://(dx\.)?doi\.org/", "", d)
    d = d.replace("\\", "/")
    # filename form -> canonical: first underscore is the 10_1016 -> 10.1016 dot,
    # but safest is to strip all separators and compare the bare token.
    d = re.sub(r"[^a-z0-9]", "", d)   # drop / . _ - spaces
    return d


def retrieved_dois(retrieved_chunks: List[Dict[str, Any]]) -> List[str]:
    out = []
    for ch in retrieved_chunks:
        meta = ch.get("metadata", {}) or {}
        out.append(normalize_doi(meta.get("doi")))
    return out


def rank_of_correct_doi(expected_doi: str, retrieved: List[Dict[str, Any]]) -> Optional[int]:
    """1-based rank of the first chunk whose DOI matches the expected DOI; None if absent."""
    target = normalize_doi(expected_doi)
    for idx, d in enumerate(retrieved_dois(retrieved), start=1):
        if d and d == target:
            return idx
    return None


# ----------------------------------------------------------------------
# Answer checking
# ----------------------------------------------------------------------
# Matches: optional sign, integer part with optional thousands separators OR a
# plain run of digits, optional decimal, optional e-exponent. Crucially the
# integer part is \d+ (not \d{1,3}) so '3263' is NOT split into 326 + 3.
_NUM_RE = re.compile(r"[-+]?\d{1,3}(?:,\d{3})+(?:\.\d+)?|[-+]?\d+(?:\.\d+)?(?:e[-+]?\d+)?")

def extract_numbers(text: str) -> List[float]:
    """Pull numeric values out of text, tolerant of thousands separators and
    scientific notation written as '5.3 x 10-4' / '5.3e-4' / '5.3 × 10⁻⁴'.

    Strategy: FIRST collapse any 'mantissa × 10^exp' form (in any notation) into
    a single e-number, THEN scan. This prevents the exponent ('-4') and base
    ('10') leaking out as separate numbers, and prevents a bare integer like
    '3263' from being mis-split by a thousands-grouping pattern."""
    if not text:
        return []
    t = str(text)

    # 1) Unicode superscripts -> ASCII so '10⁻⁴' becomes '10-4'
    superscript = str.maketrans("⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻", "0123456789+-")
    t = t.translate(superscript)

    # 2) Collapse 'B × 10^E' / 'B x 10^E' / 'B·10^E' (E possibly signed) into 'Be E'
    #    e.g. '5.3 × 10-4' -> '5.3e-4', '1.2 x 10 6' -> '1.2e6'
    t = re.sub(
        r"([-+]?\d+(?:\.\d+)?)\s*[×x·]\s*10\s*\^?\s*([-+]?\d+)",
        lambda m: f"{m.group(1)}e{m.group(2)}",
        t,
    )
    t = t.replace("E", "e")

    nums: List[float] = []
    for m in _NUM_RE.findall(t):
        s = m.replace(",", "")
        if s in ("", "-", "+", ".", "e"):
            continue
        try:
            nums.append(float(s))
        except ValueError:
            pass
    return nums


def auto_answer_match(expected: str, actual: str, qtype: str) -> Dict[str, Any]:
    """Deterministic check. For NUMBER/HARD: do the expected numbers appear in the
    answer (within tolerance)? For PROSE: substring / token-overlap heuristic."""
    expected = (expected or "").strip()
    actual = (actual or "").strip()
    if not actual:
        return {"auto_pass": False, "auto_reason": "empty answer"}

    exp_nums = extract_numbers(expected)
    if qtype in ("NUMBER", "HARD") and exp_nums:
        act_nums = extract_numbers(actual)
        matched = []
        for en in exp_nums:
            tol = max(abs(en) * 0.01, 1e-9)  # 1% relative tolerance
            hit = any(abs(en - an) <= tol for an in act_nums)
            matched.append(hit)
        passed = all(matched)
        return {
            "auto_pass": passed,
            "auto_reason": f"expected nums {exp_nums} matched={matched}",
        }

    # PROSE / fallback: case-insensitive substring, else token overlap
    el, al = expected.lower(), actual.lower()
    if el and el in al:
        return {"auto_pass": True, "auto_reason": "exact substring"}
    exp_tokens = set(re.findall(r"[a-z0-9]+", el))
    act_tokens = set(re.findall(r"[a-z0-9]+", al))
    if exp_tokens:
        overlap = len(exp_tokens & act_tokens) / len(exp_tokens)
        return {
            "auto_pass": overlap >= 0.6,
            "auto_reason": f"token overlap {overlap:.0%}",
        }
    return {"auto_pass": False, "auto_reason": "no basis for match"}


def llm_judge(llm, question: str, expected: str, actual: str) -> Dict[str, Any]:
    """Ask the loaded LLM whether the actual answer matches the expected answer.
    Returns {'judge_pass': bool, 'judge_raw': str}. Uses the same LLMWrapper model."""
    prompt = (
        "You are grading a question-answering system. Decide whether the SYSTEM ANSWER "
        "is correct, using the REFERENCE ANSWER as ground truth. Minor wording or unit "
        "formatting differences are acceptable; the facts/values must match.\n\n"
        f"QUESTION: {question}\n"
        f"REFERENCE ANSWER: {expected}\n"
        f"SYSTEM ANSWER: {actual}\n\n"
        "Reply with exactly one word on the first line: PASS or FAIL. "
        "Optionally add a brief reason on the next line."
    )
    # Reuse the wrapper's generate() but with no retrieval context.
    raw = llm.generate(prompt, [])
    verdict = raw.strip().splitlines()[0].upper() if raw.strip() else ""
    passed = verdict.startswith("PASS")
    return {"judge_pass": passed, "judge_raw": raw.strip()[:300]}


# ----------------------------------------------------------------------
# Chunk-hit: were the answer's key facts actually present in the retrieved chunks?
# The tier between paper-hit and answer-correct: "was the answer even VISIBLE?"
# Lenient on wording (tolerates paraphrase / hyphenation), strict on facts
# (numbers, distinctive terms, acronyms must appear). Heuristic; validated by hand
# against the v2 set, where it reproduced the manual chunk-hit count.
# ----------------------------------------------------------------------
_CHUNK_STOP = {
    "the", "and", "with", "for", "was", "were", "used", "uses", "using", "via", "from",
    "study", "studies", "method", "methods", "approach", "database", "functional", "unit",
    "units", "impact", "assessment", "analysis", "model", "based", "both", "main", "data",
    "this", "that", "are", "per", "over", "into", "across", "each", "their", "which",
    "considered", "corresponding", "respectively", "system", "boundary", "scope",
}


def _norm_text(s: str) -> str:
    s = str(s).lower().replace("-", "")          # protein-corrected -> proteincorrected
    s = re.sub(r"[^\w.%×]", " ", s)              # keep word chars, dots, %, x; rest -> space
    return re.sub(r"\s+", " ", s)


def _key_tokens(expected: str) -> set:
    e = _norm_text(expected)
    toks = set()
    toks |= set(re.findall(r"\d+(?:\.\d+)*", e))                       # numbers: 3.8, 10.1.0.4, 365, 1
    toks |= {w for w in re.findall(r"[a-z]{4,}", e) if w not in _CHUNK_STOP}  # words >=4, non-stop
    toks |= {a.lower() for a in re.findall(r"\b[A-Z]{3,}\d*\b", str(expected))}  # acronyms >=3: CML, FPCM
    return toks


def chunk_hit_check(expected: str, expected_doi: str, retrieved: List[Dict[str, Any]],
                    threshold: float, any_chunk: bool) -> Dict[str, Any]:
    """Did >= threshold of the expected answer's key facts appear in the retrieved
    chunks (target-paper chunks only, unless any_chunk)? Returns chunk_hit (bool|None),
    the matched fraction, and the missed tokens. None when there is no expected answer
    to check against. Target-paper filtering means a paper-miss is automatically a
    chunk-miss (no target chunks -> nothing matches)."""
    toks = _key_tokens(expected)
    if not toks:
        return {"chunk_hit": None, "chunk_hit_frac": None, "chunk_hit_missed": []}
    target = normalize_doi(expected_doi)
    if any_chunk:
        use = retrieved
    else:
        use = [c for c in retrieved
               if normalize_doi((c.get("metadata", {}) or {}).get("doi")) == target]
    blob = _norm_text(" ".join((c.get("text") or "") for c in use))
    matched = {t for t in toks if t in blob}
    frac = len(matched) / len(toks)
    return {
        "chunk_hit": frac >= threshold,
        "chunk_hit_frac": round(frac, 3),
        "chunk_hit_missed": sorted(toks - matched),
    }


# ----------------------------------------------------------------------
# Per-question result
# ----------------------------------------------------------------------
@dataclass
class QResult:
    id: str
    type: str
    question: str
    expected_answer: str
    source_doi: str
    # retrieval
    correct_rank: Optional[int] = None
    hit_at_k: bool = False
    reciprocal_rank: float = 0.0
    retrieved_dois: List[str] = field(default_factory=list)
    # chunk-hit (retrieval coverage: was the answer visible in the chunks?)
    chunk_hit: Optional[bool] = None
    chunk_hit_frac: Optional[float] = None
    chunk_hit_missed: List[str] = field(default_factory=list)
    # answer (optional)
    answer: Optional[str] = None
    auto_pass: Optional[bool] = None
    auto_reason: str = ""
    judge_pass: Optional[bool] = None
    judge_raw: str = ""
    # citation sanity (optional, only if answer generated)
    cited_correct_doi: Optional[bool] = None


# ----------------------------------------------------------------------
# Citation check: does a [n] marker in the answer point at the correct DOI?
# ----------------------------------------------------------------------
def check_citation(answer: str, retrieved: List[Dict[str, Any]], expected_doi: str) -> Optional[bool]:
    if not answer:
        return None
    markers = set()
    for _grp in re.findall(r"\[([^\]]+)\]", answer):
        markers.update(int(_n) for _n in re.findall(r"\d+", _grp))
    if not markers:
        return False
    target = normalize_doi(expected_doi)
    for n in markers:
        if 1 <= n <= len(retrieved):
            d = normalize_doi((retrieved[n - 1].get("metadata", {}) or {}).get("doi"))
            if d == target:
                return True
    return False


# ----------------------------------------------------------------------
# Main evaluation loop
# ----------------------------------------------------------------------
def run_eval(args) -> None:
    questions = load_eval_set(args.eval_set)
    print(f"Loaded {len(questions)} questions from {args.eval_set}")
    from collections import Counter
    print("  type mix:", dict(Counter(q.type for q in questions)))
    if not any(q.expected_answer for q in questions):
        if args.generate:
            # SystemExit, not sys.exit: run_eval imports sys locally further down
            raise SystemExit("ERROR: this needs the full private eval workbook (expected_answer/source_location), "
                             "which is not published. --generate scores answers against expected_answer; "
                             f"{args.eval_set} has none.")
        print("  NOTE: no expected_answer in this eval set -> chunk-hit skipped "
              "(paper-level Hit@k and MRR are unaffected)")

    llm_model = args.llm_model if args.generate else None

    # Build the retrieval backend: hybrid (default; dense + BM25) or dense.
    # Both expose the same .query(question, top_k, generate_answer) -> dict
    # interface, so the scoring loop below is identical for either.
    if args.rerank:
        try:
            import os, sys
            _root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
            if _root not in sys.path:
                sys.path.insert(0, _root)
            from rerank import build_rerank_retriever
        except ImportError:
            print("ERROR: could not import build_rerank_retriever from rerank.py.")
            print("Make sure rerank.py is alongside rag.py and sentence-transformers is installed.")
            raise
        pipeline = build_rerank_retriever(
            chroma_dir=args.chroma_dir,
            chroma_collection=args.collection,
            embedding_model=args.embedding_model,
            base=args.retriever,                 # rerank on top of --retriever (default hybrid) or dense
            bm25_dir=args.bm25_dir,
            fusion=args.fusion,
            rrf_k=args.rrf_k,
            alpha=args.alpha,
            llm_model=llm_model,
            reranker_model=args.reranker_model,
            within_paper=args.rerank_within_paper,
            candidate_k=args.candidate_k,
            batch_size=args.rerank_batch_size,
        )
        print(f"  retriever: RERANK ({args.reranker_model}) "
              f"over base={args.retriever}, candidate_k={args.candidate_k}")
    elif args.retriever == "hybrid":
        try:
            import os, sys
            _root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
            if _root not in sys.path:
                sys.path.insert(0, _root)
            from hybrid import build_hybrid_retriever
        except ImportError as exc:
            print("ERROR: could not import build_hybrid_retriever from hybrid.py.")
            print("Make sure hybrid.py is alongside rag.py and bm25s is installed.")
            raise
        pipeline = build_hybrid_retriever(
            chroma_dir=args.chroma_dir,
            chroma_collection=args.collection,
            embedding_model=args.embedding_model,
            bm25_dir=args.bm25_dir,
            llm_model=llm_model,
            rrf_k=args.rrf_k,
            candidate_k=args.candidate_k,
            fusion=args.fusion,
            alpha=args.alpha,
        )
        print(f"  retriever: HYBRID (dense + BM25, fusion={args.fusion}, "
              f"candidate_k={args.candidate_k}"
              + (f", rrf_k={args.rrf_k}" if args.fusion == 'rrf' else f", alpha={args.alpha}")
              + ")")
    else:
        try:
            import os, sys
            _root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
            if _root not in sys.path:
                sys.path.insert(0, _root)
            from rag import build_default_pipeline
        except ImportError as exc:
            print("ERROR: could not import build_default_pipeline from rag.py.")
            print("Run this script from the directory that contains rag.py.")
            raise
        pipeline = build_default_pipeline(
            chroma_dir=args.chroma_dir,
            chroma_collection=args.collection,
            embedding_model=args.embedding_model,
            llm_model=llm_model,
        )
        print("  retriever: DENSE (Chroma vectors only)")

    try:
        total_vectors = pipeline.vector_store.count()
    except Exception as e:
        print(f"ERROR: could not count vectors in Chroma collection '{args.collection}': {e}")
        sys.exit(1)
    print(f"  vector store '{args.collection}' has {total_vectors} chunks")
    if total_vectors == 0:
        print(f"ERROR: Chroma collection '{args.collection}' is empty (corpus not on this machine?).")
        sys.exit(1)

    # LLM lives on .dense for the hybrid/rerank wrappers; on the pipeline itself otherwise.
    judge_llm = getattr(pipeline, "llm", None) or getattr(getattr(pipeline, "dense", None), "llm", None)
    do_judge = args.generate and args.judge
    if do_judge and judge_llm is None:
        print("  NOTE: --judge requested but no LLM loaded; disabling judge.")
        do_judge = False

    results: List[QResult] = []
    t0 = time.time()
    for i, q in enumerate(questions, start=1):
        out = pipeline.query(
            question=q.question,
            top_k=args.top_k,
            generate_answer=args.generate,
        )
        retrieved = out.get("retrieved_chunks", [])
        rank = rank_of_correct_doi(q.source_doi, retrieved)

        r = QResult(
            id=q.id, type=q.type, question=q.question,
            expected_answer=q.expected_answer, source_doi=q.source_doi,
            correct_rank=rank,
            hit_at_k=rank is not None,
            reciprocal_rank=(1.0 / rank) if rank else 0.0,
            retrieved_dois=retrieved_dois(retrieved),
        )

        # chunk-hit runs in every mode (no LLM needed) — was the answer visible?
        ch = chunk_hit_check(q.expected_answer, q.source_doi, retrieved,
                             args.chunk_hit_threshold, args.chunk_hit_any)
        r.chunk_hit = ch["chunk_hit"]
        r.chunk_hit_frac = ch["chunk_hit_frac"]
        r.chunk_hit_missed = ch["chunk_hit_missed"]

        if args.generate:
            ans = out.get("answer") or ""
            r.answer = ans
            am = auto_answer_match(q.expected_answer, ans, q.type)
            r.auto_pass = am["auto_pass"]
            r.auto_reason = am["auto_reason"]
            r.cited_correct_doi = check_citation(ans, retrieved, q.source_doi)
            if do_judge:
                jm = llm_judge(judge_llm, q.question, q.expected_answer, ans)
                r.judge_pass = jm["judge_pass"]
                r.judge_raw = jm["judge_raw"]

        results.append(r)
        flag = "HIT " if r.hit_at_k else "miss"
        cflag = "" if r.chunk_hit is None else ("  chunk=Y" if r.chunk_hit else "  chunk=.")
        extra = ""
        if args.generate:
            ap = "P" if r.auto_pass else "."
            jp = "" if r.judge_pass is None else ("/J:P" if r.judge_pass else "/J:.")
            extra = f"  ans[{ap}{jp}]"
        print(f"  [{i:>2}/{len(questions)}] {q.type:<6} {flag} rank={r.correct_rank}{cflag}{extra}")

    elapsed = time.time() - t0
    summarize_and_save(results, args, elapsed)


def summarize_and_save(results: List[QResult], args, elapsed: float) -> None:
    n = len(results)
    hits = sum(1 for r in results if r.hit_at_k)
    mrr = sum(r.reciprocal_rank for r in results) / n if n else 0.0
    chunk_n = sum(1 for r in results if r.chunk_hit is not None)
    chunk_ok = sum(int(r.chunk_hit) for r in results if r.chunk_hit is not None)

    # by-type breakdown
    from collections import defaultdict
    by_type = defaultdict(lambda: {"n": 0, "hits": 0, "rr": 0.0,
                                   "chunk": 0, "chunk_n": 0,
                                   "auto": 0, "auto_n": 0, "judge": 0, "judge_n": 0,
                                   "cite": 0, "cite_n": 0})
    for r in results:
        b = by_type[r.type]
        b["n"] += 1
        b["hits"] += int(r.hit_at_k)
        b["rr"] += r.reciprocal_rank
        if r.chunk_hit is not None:
            b["chunk_n"] += 1; b["chunk"] += int(r.chunk_hit)
        if r.auto_pass is not None:
            b["auto_n"] += 1; b["auto"] += int(r.auto_pass)
        if r.judge_pass is not None:
            b["judge_n"] += 1; b["judge"] += int(r.judge_pass)
        if r.cited_correct_doi is not None:
            b["cite_n"] += 1; b["cite"] += int(r.cited_correct_doi)

    print("\n" + "=" * 64)
    print(f"SUMMARY   (k={args.top_k}, {n} questions, {elapsed:.1f}s)")
    print(f"  embedding model : {args.embedding_model}")
    print(f"  collection      : {args.collection}")
    print(f"  retriever       : {args.retriever}")
    print("=" * 64)
    print(f"  RETRIEVAL   Hit@{args.top_k}: {hits}/{n} = {hits/n:.1%}    MRR: {mrr:.3f}")
    if chunk_n:
        print(f"  CHUNK-HIT   {chunk_ok}/{chunk_n} = {chunk_ok/chunk_n:.1%}  "
              f"(answer's key facts present in retrieved chunks, >= {args.chunk_hit_threshold:.0%})")
    if args.generate:
        auto_n = sum(b['auto_n'] for b in by_type.values())
        auto_ok = sum(b['auto'] for b in by_type.values())
        if auto_n:
            print(f"  ANSWER auto : {auto_ok}/{auto_n} = {auto_ok/auto_n:.1%}")
        jn = sum(b['judge_n'] for b in by_type.values())
        jok = sum(b['judge'] for b in by_type.values())
        if jn:
            print(f"  ANSWER judge: {jok}/{jn} = {jok/jn:.1%}")
        cn = sum(b['cite_n'] for b in by_type.values())
        cok = sum(b['cite'] for b in by_type.values())
        if cn:
            print(f"  CITATION ok : {cok}/{cn} = {cok/cn:.1%}  (a [n] marker points at the correct DOI)")

    print("\n  by type:")
    print(f"    {'type':<18}{'n':>4}{'Hit@k':>8}{'chunk':>8}{'MRR':>7}{'auto':>7}{'judge':>7}")
    for t, b in sorted(by_type.items()):
        hit = b["hits"] / b["n"] if b["n"] else 0
        rr = b["rr"] / b["n"] if b["n"] else 0
        cu = f"{b['chunk']}/{b['chunk_n']}" if b["chunk_n"] else "-"
        au = f"{b['auto']}/{b['auto_n']}" if b["auto_n"] else "-"
        ju = f"{b['judge']}/{b['judge_n']}" if b["judge_n"] else "-"
        print(f"    {t:<18}{b['n']:>4}{hit:>7.0%}{cu:>8}{rr:>7.2f}{au:>7}{ju:>7}")

    # ---- diagnostic: where the answer leaked (only meaningful with --generate) ----
    if args.generate:
        def _ans_ok(r):
            return r.judge_pass if r.judge_pass is not None else r.auto_pass
        usable = [r.id for r in results if r.chunk_hit and not _ans_ok(r)]
        table_gap = [r.id for r in results if r.hit_at_k and r.chunk_hit is False]
        sig = "judge" if any(r.judge_pass is not None for r in results) else "auto"
        print(f"\n  --- diagnostic (answer signal: {sig}) ---")
        print(f"  chunk-hit BUT answer-wrong (retrieved & visible, LLM failed): {usable or '[]'}")
        print(f"  paper-hit BUT chunk-miss (answer not in retrieved text)     : {table_gap or '[]'}")

    # ---- save outputs ----
    os.makedirs(args.out, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    tag = f"_{args.tag}" if args.tag else ""

    json_path = os.path.join(args.out, f"results_{ts}{tag}.json")
    payload = {
        "timestamp": ts,
        "config": {
            "eval_set": args.eval_set, "top_k": args.top_k,
            "embedding_model": args.embedding_model, "collection": args.collection,
            "retriever": args.retriever, "fusion": args.fusion,
            "candidate_k": args.candidate_k, "rrf_k": args.rrf_k, "alpha": args.alpha,
            "bm25_dir": args.bm25_dir, "rerank": args.rerank,
            "reranker_model": args.reranker_model if args.rerank else None,
            "rerank_within_paper": args.rerank_within_paper,
            "chunk_hit_threshold": args.chunk_hit_threshold, "chunk_hit_any": args.chunk_hit_any,
            "generate": args.generate, "judge": args.generate and args.judge,
            "llm_model": args.llm_model if args.generate else None,
            "tag": args.tag,
        },
        "summary": {
            "n": n, "hit_at_k": hits, "hit_rate": hits / n if n else 0, "mrr": mrr,
            "chunk_hit": chunk_ok, "chunk_hit_n": chunk_n,
            "chunk_hit_rate": (chunk_ok / chunk_n) if chunk_n else None,
        },
        "results": [asdict(r) for r in results],
    }
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2, ensure_ascii=False)

    csv_path = os.path.join(args.out, f"scorecard_{ts}{tag}.csv")
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["id", "type", "hit@k", "rank", "recip_rank",
                    "chunk_hit", "chunk_hit_frac",
                    "auto_pass", "judge_pass", "cited_correct_doi",
                    "expected_answer", "answer", "question", "source_doi", "auto_reason"])
        for r in results:
            w.writerow([
                r.id, r.type, int(r.hit_at_k), r.correct_rank, f"{r.reciprocal_rank:.3f}",
                "" if r.chunk_hit is None else int(r.chunk_hit),
                "" if r.chunk_hit_frac is None else f"{r.chunk_hit_frac:.3f}",
                "" if r.auto_pass is None else int(r.auto_pass),
                "" if r.judge_pass is None else int(r.judge_pass),
                "" if r.cited_correct_doi is None else int(r.cited_correct_doi),
                r.expected_answer, (r.answer or "").replace("\n", " ")[:500],
                r.question, r.source_doi, r.auto_reason,
            ])

    print(f"\n  saved:\n    {json_path}\n    {csv_path}")
    print("\n  Tip: re-run after each change with a different --tag (e.g. --tag hybrid_k5)")
    print("       then diff the scorecards to see what moved.")


def build_argparser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="Score the LCA RAG system against an eval set.")
    p.add_argument("--eval-set", default="eval/LCA_RAG_eval_set_v3_50.xlsx",
                   help="Path to the eval set: .xlsx or .csv (default: the private v3_50 workbook, not published)")
    p.add_argument("--top-k", type=int, default=10, help="How many chunks to retrieve (default 10)")
    p.add_argument("--chroma-dir", default="chroma_db", help="Chroma persistence directory")
    p.add_argument("--collection", default="papers", help="Chroma collection name")
    p.add_argument("--embedding-model",
                   default="sentence-transformers/all-MiniLM-L6-v2",
                   help="Embedding model (match what the collection was built with!)")
    # --- retriever backend (Stage 2) ---
    p.add_argument("--retriever", choices=["dense", "hybrid"], default="hybrid",
                   help="Retrieval backend: hybrid (default, dense+BM25) or dense (Chroma only)")
    p.add_argument("--bm25-dir", default=None,
                   help="BM25 index dir (hybrid only; default bm25_index/<collection>)")
    p.add_argument("--fusion", choices=["rrf", "weighted"], default="weighted",
                   help="Hybrid fusion method (hybrid only; default weighted = validated config)")
    p.add_argument("--candidate-k", type=int, default=50,
                   help="Candidates pulled from EACH retriever before fusing (hybrid only; "
                        "default 50 = README's measured config; runtime uses 100)")
    p.add_argument("--rrf-k", type=int, default=60,
                   help="RRF constant (hybrid + fusion=rrf only; default 60)")
    p.add_argument("--alpha", type=float, default=0.4,
                   help="Dense weight for fusion=weighted (hybrid only; 0..1; default 0.4 = validated)")
    # --- reranking (Stage 2, second component) ---
    p.add_argument("--rerank", action="store_true",
                   help="Add a cross-encoder reranker on top of the base retriever")
    p.add_argument("--reranker-model", default="BAAI/bge-reranker-v2-m3",
                   help="Cross-encoder model for reranking")
    p.add_argument("--rerank-within-paper", action="store_true",
                   help="Rerank chunks only WITHIN each paper; freeze paper order from base "
                        "(protects Hit@k/MRR; only chunk selection changes)")
    p.add_argument("--rerank-batch-size", type=int, default=32,
                   help="Cross-encoder batch size (--rerank only; default 32)")
    # --- chunk-hit (Stage 2 coverage metric) ---
    p.add_argument("--chunk-hit-threshold", type=float, default=0.6,
                   help="Min fraction of expected-answer key facts that must appear in the "
                        "retrieved chunks to count as a chunk-hit (default 0.6)")
    p.add_argument("--chunk-hit-any", action="store_true",
                   help="Count key facts in ANY retrieved chunk, not only target-paper chunks "
                        "(looser; risks crediting sibling papers)")
    # --- generation / judging ---
    p.add_argument("--generate", action="store_true",
                   help="Also generate answers and score them (slower, needs --llm-model)")
    p.add_argument("--llm-model", default="Qwen/Qwen2.5-7B-Instruct",
                   help="Model tag for answer generation (used only with --generate; "
                        "for the Ollama backend pass e.g. qwen2.5:7b-instruct-q8_0)")
    p.add_argument("--judge", action="store_true",
                   help="Use the LLM to grade answers (only with --generate)")
    p.add_argument("--out", default="eval_runs", help="Output directory for results")
    p.add_argument("--tag", default="", help="Label for this run (appears in filenames)")
    return p


if __name__ == "__main__":
    args = build_argparser().parse_args()
    if args.judge and not args.generate:
        print("Note: --judge has no effect without --generate. Enabling retrieval-only.")
    run_eval(args)