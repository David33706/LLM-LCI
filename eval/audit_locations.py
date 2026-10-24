#!/usr/bin/env python3
"""
audit_locations.py — SECTION-HIT audit: computed `location` vs gold source_location.

For every tagged-gold yes-row in a --dump-answers file, compare the Python-derived
section against the eval set's gold location. Completes the primary (localization)
metric at section level. Offline: no GPU, no LLM.

Needs the gold source_location column, which only the full eval workbook has. That
workbook is kept in the author's private archive (not published: it contains verbatim
paper text or licensed data); a .csv questions file cannot be used here.

Run from the repo root:
  python eval/audit_locations.py                  # newest dump, default eval file
  python eval/audit_locations.py --dump eval_runs/answers_X.xlsx --qfile path/to/eval_workbook.xlsx
"""
import argparse, glob, os, re, sys
import openpyxl

PRIVATE_ERR = ("ERROR: this needs the full private eval workbook (expected_answer/source_location), "
               "which is not published")

def norm_tokens(s: str) -> set:
    s = re.sub(r"\bP\d+\b", " ", str(s or ""))        # page markers: "P3"
    s = re.sub(r"\b[\d.]+\b", " ", s)                 # numbering: "2.2.1.2."
    toks = {w.lower() for w in re.findall(r"[A-Za-z][A-Za-z0-9]{2,}", s)}
    return {t[:-1] if len(t) > 4 and t.endswith("s") else t for t in toks}  # crude plural

def is_match(gold: str, derived: str) -> bool:
    g, d = norm_tokens(gold), norm_tokens(derived)
    if not g or not d:
        return False
    inter = g & d
    return g <= d or d <= g or (len(inter) / len(g)) >= 0.6

def load_gold_locations(qfile):
    wb = openpyxl.load_workbook(qfile, data_only=True)
    ws = wb["Questions"] if "Questions" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h).strip().lower() if h else "" for h in rows[0]]
    loc_col = next((c for c in ("source_location", "location") if c in hdr), None)
    if loc_col is None:
        sys.exit(f"{PRIVATE_ERR}. {qfile} has no source_location column.")
    id_col = "id" if "id" in hdr else None
    out = {}
    for i, r in enumerate(rows[1:], 1):
        if r is None or all(v is None for v in r):
            continue
        row = {hdr[j]: (r[j] if j < len(r) else None) for j in range(len(hdr))}
        qid = str(row.get(id_col)).strip() if id_col else str(i)   # authoring file: row order = qid
        loc = row.get(loc_col) if loc_col else None
        out[qid] = str(loc).strip() if loc not in (None, "", "None") else ""
    return out

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dump", default=None, help="answers_*.xlsx (default: newest in eval_runs/)")
    ap.add_argument("--qfile", default="eval/LCA_RAG_eval_set_v3_50.xlsx",
                    help="full eval workbook with source_location (private, not published)")
    args = ap.parse_args()
    if args.qfile.lower().endswith(".csv") or not os.path.exists(args.qfile):
        sys.exit(f"{PRIVATE_ERR}. Got --qfile {args.qfile}"
                 + (" (file not found)." if not os.path.exists(args.qfile) else
                    " (a .csv has no source_location)."))
    dump = args.dump or (sorted(glob.glob("eval_runs/answers_*.xlsx")) or [None])[-1]
    if not dump:
        sys.exit("no answers_*.xlsx found; pass --dump")
    if not os.path.exists(dump):
        sys.exit(f"ERROR: answers dump not found: {dump}")

    gold_loc = load_gold_locations(args.qfile)
    ws = openpyxl.load_workbook(dump, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    hdr = list(rows[0])
    data = [dict(zip(hdr, r)) for r in rows[1:] if r and r[0] is not None]
    gold_rows = [d for d in data if str(d.get("is_tagged_gold", "")).strip() == "YES"]

    match, mismatch, no_gold = [], [], []
    for d in gold_rows:
        qid = str(d["qid"])
        g = gold_loc.get(qid, "")
        if not g:
            no_gold.append((qid, d["location"]))
        elif is_match(g, d["location"]):
            match.append((qid, g, d["location"]))
        else:
            mismatch.append((qid, g, d["location"]))

    print(f"dump: {os.path.basename(dump)} | gold yes-rows: {len(gold_rows)}")
    print(f"\nSECTION-HIT: {len(match)}/{len(match)+len(mismatch)} scoreable "
          f"({len(no_gold)} rows have no gold location)\n")
    print("MATCHES:")
    for qid, g, d in match:
        print(f"  Q{qid:<3} gold '{g[:42]:<42}' ~ derived '{d[:42]}'")
    print("\nMISMATCHES:")
    for qid, g, d in mismatch:
        print(f"  Q{qid:<3} gold '{g[:42]:<42}' vs derived '{d[:42]}'")
    print("\nNO_GOLD (source_location missing in eval set):")
    print("  qids:", ", ".join(q for q, _ in no_gold) or "none")

if __name__ == "__main__":
    main()