#!/usr/bin/env python3
"""
apply_valid_dois.py — write Prof. Lam's validated alternatives (2026-07-16 reply)
into the eval workbook as a `valid_dois` column.

ALREADY APPLIED. The full eval workbook already holds these values: 12 qids (1, 11,
14, 16, 20, 21, 28, 31, 33, 34, 35, 38), 22 (qid, DOI) pairs, 21 distinct DOIs.
Check with: python eval/validate_eval_set.py

The full workbook and the expert's validation file are kept in the author's private
archive (not published: they contain verbatim paper text or licensed data). Without
the workbook this script exits with a message and writes nothing.

This script REWRITES a frozen workbook (a load-bearing wall, RUNBOOK §8). Do NOT
re-run it without approval. Idempotent: re-running overwrites the same cells.

If ever needed, run from the repo root:
  python eval/apply_valid_dois.py
"""
import os
import sys

import openpyxl

PATH = "eval/LCA_RAG_eval_set_v3_50.xlsx"   # the full private workbook (not published)

# Prof. Lam's Y-marked studies, as {qid: "doi; doi"}. They are part of the private eval
# set and are not published; fill this in locally from his validation workbook.
VALID = {}

if not VALID:
    sys.exit("VALID is empty: the expert-validated (qid, DOI) pairs are part of the private "
             "eval set and are not published.")
if not os.path.exists(PATH):
    sys.exit(f"{PATH} not found: this script edits the full private eval workbook, which is "
             "not published.")

wb = openpyxl.load_workbook(PATH)
ws = wb["Questions"] if "Questions" in wb.sheetnames else wb[wb.sheetnames[0]]
hdr = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]

if "valid_dois" in hdr:
    vcol = hdr.index("valid_dois") + 1
else:
    vcol = len(hdr) + 1
    ws.cell(row=1, column=vcol, value="valid_dois")

idcol = hdr.index("id") + 1
applied = []
for row in ws.iter_rows(min_row=2):
    qid = str(row[idcol - 1].value).strip() if row[idcol - 1].value is not None else ""
    if qid in VALID:
        ws.cell(row=row[0].row, column=vcol, value=VALID[qid])
        applied.append(qid)

wb.save(PATH)
print(f"valid_dois written for qids: {applied} -> {PATH}")
print("verify:")
for row in ws.iter_rows(min_row=2):
    qid = str(row[idcol - 1].value).strip() if row[idcol - 1].value is not None else ""
    if qid in VALID:
        print(f"  Q{qid}: {ws.cell(row=row[0].row, column=vcol).value[:70]}...")