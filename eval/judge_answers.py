#!/usr/bin/env python3
"""
judge_answers.py — ANSWER-ACCURACY judge: quote vs expected_answer (reference-based).

For each tagged-gold yes-row in a --dump-answers file, ask an independent LLM call:
"does the verbatim quote state the same assumption as the reference answer?"
The judge sees ONLY (STUDY TITLE, question, expected_answer, quote) — not the paper,
not the first model's verdict. The STUDY TITLE line comes from --wos
(paper_list2_2025.csv) and is part of the measurement: a missing/empty titles file is
an ERROR unless --no-titles is passed deliberately. Produces the secondary metric:
answer accuracy X/N.

--target expected needs expected_answer, which only the full eval workbook has. That
workbook is kept in the author's private archive (not published: it contains verbatim
paper text or licensed data); a .csv questions file cannot be used here.

Run from the repo root (needs the vLLM server):
  # main: reference-based, gold rows only (dump from step3_orchestrate.py --dump-answers)
  python eval/judge_answers.py --dump eval_runs/answers_<ts>.xlsx

  # calibration: reference-FREE (quote vs question), all rows of given qids;
  # verdicts must reproduce Prof. Lam's 9 human labels on qids 1,31,32
  python eval/judge_answers.py --dump ... --target question --qids 1,31,32

The output judged_<target>_*.xlsx is never overwritten unless --overwrite is given.
"""
import argparse, csv, glob, json, os, re, sys, time
import openpyxl
import requests

PRIVATE_ERR = ("ERROR: this needs the full private eval workbook (expected_answer/source_location), "
               "which is not published")


def load_titles(wos_csv):
    titles = {}
    try:
        with open(wos_csv, newline="", encoding="utf-8-sig") as f:
            for row in csv.DictReader(f):
                d = re.sub(r"[^a-z0-9]", "", str(row.get("DOI") or "").lower())
                if d:
                    titles[d] = (row.get("Article Title") or "").strip()
    except FileNotFoundError:
        pass
    return titles

JUDGE_PROMPT_REF = """You compare a REFERENCE answer with a QUOTE from a scientific paper.
Decide if the quote states the SAME assumption/value/practice as the reference.
Rules: paraphrase is fine; numbers and units must agree with the reference (ranges
that contain the reference value count); extra detail in the quote is fine; if the
quote does not state the referenced fact at all, it is not a match.
Respond ONLY with JSON: {"match": true/false, "reason": "one short sentence"}"""

JUDGE_PROMPT_Q = """You check whether a QUOTE from ONE scientific study answers a QUESTION about
what LCA studies typically/commonly assume. You are given the STUDY TITLE for context.
It COUNTS as an answer if the quote states the study's own adopted value, definition,
or named choice for the asked quantity — OR reports a specific NUMERIC rate/value for
the asked quantity as routinely achieved or cited from the literature.
The quote does NOT need to claim the value is typical or common. It does NOT need to
repeat the technology name: if the study (per its title) concerns the asked process
family, statements about "the capture level" / "the process" count for that technology.
It is NOT an answer if it is merely on-topic without stating the asked quantity, or if
it is hypothetical or illustrative ("could be", "for example"). Review-style statements
that list or define the options that exist or are commonly used in general do NOT
count — naming several options is not answering; the study must itself adopt or
numerically report one.
Respond ONLY with JSON: {"match": true/false, "reason": "one short sentence"}"""


def load_expected(qfile):
    wb = openpyxl.load_workbook(qfile, data_only=True)
    ws = wb["Questions"] if "Questions" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h).strip().lower() if h else "" for h in rows[0]]
    ans_col = next((c for c in ("expected_answer", "answer") if c in hdr), None)
    if ans_col is None:
        sys.exit(f"{PRIVATE_ERR}. {qfile} has no expected_answer column.")
    id_col = "id" if "id" in hdr else None
    out = {}
    for i, r in enumerate(rows[1:], 1):
        if r is None or all(v is None for v in r):
            continue
        row = {hdr[j]: (r[j] if j < len(r) else None) for j in range(len(hdr))}
        qid = str(row.get(id_col)).strip() if id_col else str(i)
        out[qid] = str(row.get(ans_col) or "").strip() if ans_col else ""
    return out


def parse_match(v):
    """Strict verdict: True/False as is; "true"/"yes"/"match" -> True; "false"/"no" -> False;
    missing key or anything else -> None (ERR). bool("false") would be True."""
    if isinstance(v, bool):
        return v
    if isinstance(v, str):
        t = v.strip().lower()
        if t in ("true", "yes", "match"):
            return True
        if t in ("false", "no"):
            return False
    return None


def judge_one(url, model, system, user, max_retries=2, no_think=False):
    payload = {"model": model,
               "messages": [{"role": "system", "content": system},
                            {"role": "user", "content": user}],
               "temperature": 0, "max_tokens": 200,
               "response_format": {"type": "json_object"}}

    if no_think:
        payload["chat_template_kwargs"] = {"enable_thinking": False}

    for _ in range(max_retries + 1):
        try:
            r = requests.post(url, json=payload, timeout=120)
            r.raise_for_status()
            content = r.json()["choices"][0]["message"]["content"]
            s, e = content.find("{"), content.rfind("}")
            parsed = json.loads(content[s:e + 1])
            verdict = parse_match(parsed.get("match"))
            if verdict is None:
                return None, f"judge_error: bad 'match' value {parsed.get('match')!r}"[:200]
            return verdict, str(parsed.get("reason", ""))[:200]
        except Exception as ex:
            err = str(ex)
            time.sleep(2)
    return None, f"judge_error: {err[:120]}"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dump", default=None)
    ap.add_argument("--qfile", default="eval/LCA_RAG_eval_set_v3_50.xlsx",
                    help="full eval workbook with expected_answer (private, not published)")
    ap.add_argument("--target", choices=["expected", "question"], default="expected",
                    help="expected: quote vs expected_answer, GOLD rows (main metric). "
                         "question: quote vs question, ALL rows (calibration vs human labels).")
    ap.add_argument("--qids", default=None, help="comma-separated qid filter, e.g. 1,31,32")
    ap.add_argument("--wos", default="paper_list2_2025.csv", help="DOI->title map (judge context)")
    ap.add_argument("--no-titles", action="store_true",
                    help="deliberately judge WITHOUT the STUDY TITLE line (changes the judge's input)")
    ap.add_argument("--overwrite", action="store_true",
                    help="allow replacing an existing judged_<target>_*.xlsx")
    ap.add_argument("--url", default="http://localhost:3000/v1/chat/completions")
    ap.add_argument("--model", default="qwen2.5-32b-awq")
    ap.add_argument("--no-think", action="store_true",
                    help="Qwen3+ hybrid models: disable thinking via chat_template_kwargs")
    args = ap.parse_args()

    if args.target == "expected" and (args.qfile.lower().endswith(".csv")
                                      or not os.path.exists(args.qfile)):
        sys.exit(f"{PRIVATE_ERR}. Got --qfile {args.qfile}"
                 + (" (file not found)." if not os.path.exists(args.qfile) else
                    " (a .csv has no expected_answer)."))
    dump = args.dump or (sorted(glob.glob("eval_runs/answers_*.xlsx")) or [None])[-1]
    if not dump:
        sys.exit("no answers_*.xlsx found; pass --dump")
    if not os.path.exists(dump):
        sys.exit(f"ERROR: answers dump not found: {dump}")
    out = dump.replace("answers_", f"judged_{args.target}_")
    if os.path.abspath(out) == os.path.abspath(dump):
        sys.exit(f"ERROR: output path equals --dump ({dump}); the dump name must contain 'answers_'")
    if os.path.exists(out) and not args.overwrite:
        sys.exit(f"ERROR: {out} already exists; pass --overwrite to replace it")
    expected = load_expected(args.qfile) if args.target == "expected" else {}
    titles ={} if args.no_titles else load_titles(args.wos)
    if not args.no_titles and not titles:
        print(f"ERROR: titles file {args.wos} missing/empty; the judge's STUDY TITLE line is part "
              f"of the measurement (pass --wos paper_list2_2025.csv from the repo root)")
        sys.exit(1)
    if args.no_titles:
        print("WARNING: --no-titles: judging WITHOUT the STUDY TITLE line (not comparable to titled runs)")
    qid_filter = {q.strip() for q in args.qids.split(",")} if args.qids else None

    ws = openpyxl.load_workbook(dump, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    hdr = list(rows[0])
    data = [dict(zip(hdr, r)) for r in rows[1:] if r and r[0] is not None]

    if args.target == "expected":
        pool = [d for d in data if str(d.get("is_tagged_gold", "")).strip() == "YES"
                and expected.get(str(d["qid"]), "")]
        system = JUDGE_PROMPT_REF
    else:
        pool = data
        system = JUDGE_PROMPT_Q
    if qid_filter:
        pool = [d for d in pool if str(d["qid"]) in qid_filter]

    print(f"judging {len(pool)} rows | target={args.target} | model={args.model}")
    results, n_match = [], 0
    for d in pool:
        qid = str(d["qid"])
        t = titles.get(re.sub(r"[^a-z0-9]", "", str(d["doi"]).lower()), "")
        title_line = f"STUDY TITLE: {t}\n\n" if t else ""
        if args.target == "expected":
            user = (f"{title_line}QUESTION: {d['question']}\n\nREFERENCE ANSWER: {expected[qid]}\n\n"
                    f"QUOTE FROM THE PAPER: {d['supporting_quote']}")
        else:
            user = (f"{title_line}QUESTION: {d['question']}\n\n"
                    f"QUOTE FROM THE PAPER: {d['supporting_quote']}")
        verdict, reason = judge_one(args.url, args.model, system, user, no_think=args.no_think)
        mark = {True: "MATCH", False: "no   ", None: "ERR  "}[verdict]
        if verdict:
            n_match += 1
        print(f"  Q{qid:<3} [{mark}] {str(d['doi'])[:32]:<32} {reason[:70]}")
        results.append({**d, "judge_match": mark.strip(), "judge_reason": reason})

    n_ok = sum(1 for r in results if r["judge_match"] in ("MATCH", "no"))
    print(f"\nANSWER-ACCURACY ({args.target}): {n_match}/{n_ok} judged "
          f"({len(results)-n_ok} errors)")

    wb_o = openpyxl.Workbook(); ws_o = wb_o.active
    cols = list(hdr) + ["judge_match", "judge_reason"]
    ws_o.append(cols)
    for r in results:
        ws_o.append([r.get(c, "") for c in cols])
    wb_o.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()