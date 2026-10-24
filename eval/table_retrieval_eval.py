#!/usr/bin/env python3
"""
table_retrieval_eval.py - the TABLE-LEVEL retrieval metric (the librarian metric).

WHY THIS EXISTS
---------------
evaluate.py scores retrieval at the PAPER level (Hit@k on DOI) and chunk-hit by
token-overlap with the answer key's DESCRIPTION. Neither answers what the
professor's scope poses for TASK1: "did the system retrieve the RIGHT TABLE?"
This measures exactly that, and prints the retrieved SECTION per question so you
can see table-vs-prose (which evaluate.py's JSON does not log).

STANDALONE: imports rag.py/hybrid.py from the repo root (it puts the root on sys.path,
as evaluate.py and retrieval_speed.py do; bm25_eval.py imports neither); does NOT
modify evaluate.py, so the working harness is at zero risk.

For each question whose expected_answer names a table ("... in Table N"):
  correct_table_rank : rank of first retrieved chunk from the SOURCE doi whose
                       section_title is that exact table (e.g. "Table 2")
  source_table_rank  : rank of first retrieved chunk from the source doi that is
                       ANY table (looser)
  paper_rank         : rank of first chunk from the source doi (paper-level)
Then table-Hit@k / table-MRR over those questions. Questions that don't name a
table (most TASK2_ASSUMPTION) show n/a for the table metric but still print
paper_rank + retrieved sections.

RUN FROM THE REPO ROOT (v2 is the only eval set whose answers name a "Table N"):
  python eval/table_retrieval_eval.py --eval-set path/to/eval_set_v2.xlsx \
      --retriever hybrid --fusion weighted --alpha 0.4 --bm25-dir bm25_index/papers
The v2 workbook (with expected_answer) is kept in the author's private archive (not
published: it contains verbatim paper text or licensed data). A questions file without
expected_answer makes this script exit with an ERROR.
"""
import argparse, csv, json, os, re, sys, time
from pathlib import Path

_HERE = Path(__file__).resolve().parent
_ROOT = _HERE.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))


def normalize_doi(doi):
    if not doi:
        return ""
    d = str(doi).strip().lower()
    d = re.sub(r"^https?://(dx\.)?doi\.org/", "", d)
    return re.sub(r"[^a-z0-9]", "", d)


def norm_label(s):
    return re.sub(r"[^a-z0-9]", "", (s or "").lower())


TABLE_RE = re.compile(r"\btable\s+(s?[0-9]+[a-z]?|[ivxlc]+|[a-z])\b", re.I)


def expected_table_label(expected_answer):
    """Parse 'Table N' from the answer key -> normalized like 'table2'; None if absent."""
    m = TABLE_RE.search(expected_answer or "")
    return ("table" + norm_label(m.group(1))) if m else None


def load_eval_set(path):
    """Return list of {id,type,question,expected_answer,source_doi} (xlsx 'Questions' or csv)."""
    rows = []
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True)
        ws = wb["Questions"] if "Questions" in wb.sheetnames else wb[wb.sheetnames[0]]
        header = [str(c.value).strip().lower() if c.value is not None else "" for c in ws[1]]
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r is None or all(v is None for v in r):
                continue
            rows.append({header[i]: (r[i] if i < len(r) else None) for i in range(len(header))})
    elif ext == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            for row in csv.DictReader(f):
                rows.append({(k or "").strip().lower(): v for k, v in row.items()})
    else:
        raise ValueError(f"Unsupported eval-set format: {ext}")
    out = []
    for i, row in enumerate(rows, start=1):
        q = (row.get("question") or "").strip()
        doi = (row.get("source_doi") or "").strip()
        if not q or not doi:
            continue
        out.append({
            "id": str(row.get("id") or i).strip(),
            "type": (row.get("type") or "").strip().upper() or "UNSPEC",
            "question": q,
            "expected_answer": (row.get("expected_answer") or "").strip(),
            "source_doi": doi,
        })
    if not out:
        raise ValueError(f"No usable questions in {path}")
    return out


def is_table_chunk(meta):
    st = meta.get("section_title") or ""
    return norm_label(st).startswith("table") or meta.get("section_type") == "table"


def build_retriever(args):
    if args.retriever == "hybrid":
        from hybrid import build_hybrid_retriever
        return build_hybrid_retriever(
            chroma_dir=args.chroma_dir, chroma_collection=args.collection,
            embedding_model=args.embedding_model, bm25_dir=args.bm25_dir,
            fusion=args.fusion, alpha=args.alpha,
            candidate_k=args.candidate_k, rrf_k=args.rrf_k)
    from rag import build_default_pipeline
    return build_default_pipeline(
        chroma_dir=args.chroma_dir, chroma_collection=args.collection,
        embedding_model=args.embedding_model, llm_model=None)


def main():
    ap = argparse.ArgumentParser(description="Table-level retrieval metric for the RAG system.")
    ap.add_argument("--eval-set", required=True)
    ap.add_argument("--top-k", type=int, default=10)
    ap.add_argument("--chroma-dir", default="chroma_db")
    ap.add_argument("--collection", default="papers")
    ap.add_argument("--embedding-model", default="sentence-transformers/all-MiniLM-L6-v2")
    ap.add_argument("--retriever", choices=["dense", "hybrid"], default="hybrid")
    ap.add_argument("--bm25-dir", default="bm25_index/papers")
    ap.add_argument("--fusion", choices=["rrf", "weighted"], default="weighted")
    ap.add_argument("--alpha", type=float, default=0.4)
    ap.add_argument("--candidate-k", type=int, default=100)
    ap.add_argument("--rrf-k", type=int, default=60)
    ap.add_argument("--out-dir", default="eval/eval_runs")
    ap.add_argument("--tag", default="table_retrieval")
    args = ap.parse_args()

    private_err = ("ERROR: this needs the full private eval workbook (expected_answer/source_location), "
                   "which is not published (here: the v2 workbook, whose answers name tables).")
    if not os.path.exists(args.eval_set):
        sys.exit(f"{private_err} Not found: {args.eval_set}")
    questions = load_eval_set(args.eval_set)
    if not any(q["expected_answer"] for q in questions):
        sys.exit(f"{private_err} {args.eval_set} has no expected_answer.")
    from collections import Counter
    print(f"Loaded {len(questions)} questions  type mix: {dict(Counter(q['type'] for q in questions))}")
    print(f"retriever={args.retriever} fusion={args.fusion} alpha={args.alpha} k={args.top_k}\n")

    ret = build_retriever(args)
    rows = []
    t0 = time.time()
    for q in questions:
        out = ret.query(q["question"], top_k=args.top_k, generate_answer=False)
        chunks = out.get("retrieved_chunks", [])
        src = normalize_doi(q["source_doi"])
        exp_tbl = expected_table_label(q["expected_answer"])
        paper_rank = source_table_rank = correct_table_rank = None
        seen = []
        for rank, ch in enumerate(chunks, 1):
            m = ch.get("metadata", {}) or {}
            d = normalize_doi(m.get("doi"))
            st = m.get("section_title") or "?"
            same = (d == src)
            seen.append((same, st))
            if same:
                if paper_rank is None:
                    paper_rank = rank
                if is_table_chunk(m):
                    if source_table_rank is None:
                        source_table_rank = rank
                    if exp_tbl and norm_label(st) == exp_tbl and correct_table_rank is None:
                        correct_table_rank = rank
        rows.append({
            "id": q["id"], "type": q["type"], "expected_table": exp_tbl,
            "paper_rank": paper_rank, "source_table_rank": source_table_rank,
            "correct_table_rank": correct_table_rank,
            "top_sections": seen[:args.top_k],
        })
        # per-question line: ranks + the top-k sections (★ = source paper)
        secs = "  ".join(("★" if s else " ") + st for s, st in seen[:args.top_k])
        et = exp_tbl or "n/a"
        print(f"Q{q['id']:<3} {q['type'][:8]:<8} paper={str(paper_rank):>4} "
              f"srcTbl={str(source_table_rank):>4} correctTbl={str(correct_table_rank):>4} "
              f"[want {et:<8}] | {secs}")

    elapsed = time.time() - t0

    # ---- summaries ----
    n = len(rows)
    paper_hits = sum(1 for r in rows if r["paper_rank"] and r["paper_rank"] <= args.top_k)
    tbl_rows = [r for r in rows if r["expected_table"]]
    nt = len(tbl_rows)
    correct_hits = sum(1 for r in tbl_rows if r["correct_table_rank"] and r["correct_table_rank"] <= args.top_k)
    anytbl_hits = sum(1 for r in tbl_rows if r["source_table_rank"] and r["source_table_rank"] <= args.top_k)
    tbl_mrr = (sum((1.0 / r["correct_table_rank"]) if r["correct_table_rank"] else 0.0 for r in tbl_rows) / nt) if nt else 0.0

    print("\n" + "=" * 66)
    print(f"TABLE-LEVEL RETRIEVAL   (k={args.top_k}, {n} questions, {elapsed:.1f}s)")
    print("=" * 66)
    print(f"  paper-level Hit@{args.top_k} (reference)............ {paper_hits}/{n} = {paper_hits/n:.1%}")
    print(f"  questions whose answer names a table............ {nt}")
    print(f"  RIGHT-table Hit@{args.top_k} (the librarian metric).. {correct_hits}/{nt} = {(correct_hits/nt if nt else 0):.1%}")
    print(f"  any-table-from-source Hit@{args.top_k} (looser)...... {anytbl_hits}/{nt} = {(anytbl_hits/nt if nt else 0):.1%}")
    print(f"  RIGHT-table MRR................................. {tbl_mrr:.3f}")

    # by type
    print("\n  by type (table questions only):")
    bt = {}
    for r in tbl_rows:
        b = bt.setdefault(r["type"], [0, 0])
        b[0] += 1
        b[1] += int(bool(r["correct_table_rank"] and r["correct_table_rank"] <= args.top_k))
    for t, (cnt, hh) in sorted(bt.items()):
        print(f"    {t:<18} {hh}/{cnt} right-table")

    # ---- save ----
    os.makedirs(args.out_dir, exist_ok=True)
    ts = time.strftime("%Y%m%d_%H%M%S")
    csv_path = os.path.join(args.out_dir, f"table_scorecard_{ts}_{args.tag}.csv")
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["id", "type", "expected_table", "paper_rank", "source_table_rank",
                    "correct_table_rank", "right_table_hit@k", "top_sections"])
        for r in rows:
            hit = int(bool(r["correct_table_rank"] and r["correct_table_rank"] <= args.top_k)) if r["expected_table"] else ""
            secs = " | ".join(("*" if s else "") + st for s, st in r["top_sections"])
            w.writerow([r["id"], r["type"], r["expected_table"] or "", r["paper_rank"],
                        r["source_table_rank"], r["correct_table_rank"], hit, secs])
    print(f"\n  saved: {csv_path}")


if __name__ == "__main__":
    main()