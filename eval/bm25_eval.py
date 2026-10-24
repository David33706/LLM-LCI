#!/usr/bin/env python3
"""
bm25_eval.py - Pure BM25 (keyword-only) retrieval evaluation.

WHY THIS EXISTS
---------------
A keyword-only baseline that needs no Chroma store and no embedder. (For the
record: weighted fusion in hybrid.py IS pure at its endpoints when
top_k <= candidate_k — alpha=1 reproduces the dense ranking and alpha=0 the BM25
ranking, because the other list's chunks score 0 and sort below. The June result
where alpha=0.0 and alpha=1.0 came out identical was caused by an EMPTY Chroma
store, not by the fusion.) This script answers "how much does semantic
understanding add on its own?" from the BM25 ranked list in ISOLATION.

This script does exactly that and nothing else. It does NOT import or touch
hybrid.py, rag.py, or the Chroma store, so the working hybrid path is at zero
risk. It loads the SAME BM25 index hybrid.py uses (built by ingest/build_bm25_index.py)
and the SAME eval set evaluate.py uses, ranks each question by BM25 alone, and
scores Hit@k + MRR with the SAME rank semantics as evaluate.py:
  - rank = position (1-indexed) of the first retrieved chunk whose DOI matches
    the source DOI; None if the source paper is absent from the top-k chunks
  - DOIs compared after stripping to lowercase alphanumerics
so the keyword-only number is directly comparable to the dense and hybrid runs.

RUN FROM THE REPO ROOT so bm25_index/papers resolves cleanly:
    python eval/bm25_eval.py --eval-set my_questions.csv --bm25-dir bm25_index/papers
"""

import argparse
import json
import re
import sys
import time
from collections import Counter
from pathlib import Path

# same message as step3_orchestrate.EVAL_SET_MISSING (this script does not import step3)
EVAL_SET_MISSING = ("ERROR: eval set not found: {}. The project's 50-question eval set is not published; "
                    "pass --eval-set with your own .xlsx (sheet \"Questions\") or .csv with columns "
                    "id,type,question,source_doi (optional: valid_dois, expected_answer, source_location).")


def normalize_doi(doi):
    """Lowercase + strip everything except a-z0-9 (matches evaluate.py)."""
    if not doi:
        return ""
    return re.sub(r"[^a-z0-9]", "", str(doi).lower())


def load_questions(path):
    """Read the .csv or the xlsx 'Questions' sheet -> list of {type, question, source_doi}."""
    if not Path(path).exists():
        sys.exit(EVAL_SET_MISSING.format(path))
    if str(path).lower().endswith(".csv"):
        import csv
        with open(path, newline="", encoding="utf-8-sig") as f:
            rows = [tuple(v if v.strip() else None for v in r) for r in csv.reader(f)]
    else:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb["Questions"] if "Questions" in wb.sheetnames else wb.active
        rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]

    def col(*names):
        for n in names:
            if n in header:
                return header.index(n)
        return None

    i_type = col("type", "task_type", "category")
    i_q = col("question", "query")
    i_doi = col("source_doi", "doi", "source doi")
    if i_q is None or i_doi is None:
        raise ValueError(
            f"Could not find 'question'/'source_doi' columns. Headers seen: {header}"
        )

    out = []
    for r in rows[1:]:
        if r is None or i_q >= len(r) or r[i_q] is None:
            continue
        q_type = "UNTYPED"
        if i_type is not None and i_type < len(r) and r[i_type] is not None:
            q_type = str(r[i_type]).strip()
        src = ""
        if i_doi < len(r) and r[i_doi] is not None:
            src = str(r[i_doi]).strip()
        out.append({"type": q_type, "question": str(r[i_q]).strip(), "source_doi": src})
    return out


class BM25Only:
    """Loads the BM25 index exactly as hybrid.py does; ranks by BM25 alone."""

    def __init__(self, bm25_dir):
        import bm25s

        self.bm25_dir = Path(bm25_dir)
        index_path = self.bm25_dir / "bm25s"
        if not index_path.exists():
            raise FileNotFoundError(
                f"BM25 index not found at {index_path}. "
                f"Build it with build_bm25_index.py, or fix --bm25-dir."
            )
        self._bm25 = bm25s.BM25.load(str(index_path), load_corpus=False)
        self._tokenize = bm25s.tokenize
        with open(self.bm25_dir / "chunks.json", encoding="utf-8") as f:
            side = json.load(f)
        self._dois = side["dois"]

    def ranked_dois(self, question, k):
        """DOIs of the top-k BM25 chunks, in rank order (same call as hybrid.py)."""
        q_tokens = self._tokenize(question, show_progress=False)
        results, _scores = self._bm25.retrieve(q_tokens, k=k, show_progress=False)
        row = results[0]  # single query
        dois = []
        for j in range(len(row)):
            idx = int(row[j])
            dois.append(self._dois[idx] if idx < len(self._dois) else None)
        return dois


def rank_of_correct(source_doi, retrieved_dois):
    """1-indexed rank of first retrieved DOI matching source; None if absent."""
    target = normalize_doi(source_doi)
    for i, d in enumerate(retrieved_dois, start=1):
        if normalize_doi(d) == target:
            return i
    return None


def main():
    ap = argparse.ArgumentParser(description="Pure BM25 keyword-only retrieval eval")
    ap.add_argument("--eval-set", default="eval/LCA_RAG_eval_set_v3_50.xlsx",
                    help=".xlsx or .csv (default: the private v3_50 workbook, not published)")
    ap.add_argument("--bm25-dir", default="bm25_index/papers")
    ap.add_argument("--top-k", type=int, default=10)
    ap.add_argument("--tag", default="keyword_only_pure")
    ap.add_argument("--out-dir", default="eval/eval_runs")
    args = ap.parse_args()

    questions = load_questions(args.eval_set)
    print(f"Loaded {len(questions)} questions from {args.eval_set}")
    print(f"  type mix: {dict(Counter(q['type'] for q in questions))}")
    print("  retriever: BM25-ONLY (pure keyword, no fusion, no dense)")

    bm25 = BM25Only(args.bm25_dir)

    t0 = time.time()
    rows = []
    for i, q in enumerate(questions, start=1):
        dois = bm25.ranked_dois(q["question"], k=args.top_k)
        rank = rank_of_correct(q["source_doi"], dois)
        hit = rank is not None
        rr = (1.0 / rank) if rank else 0.0
        rows.append({"type": q["type"], "rank": rank, "hit": hit, "rr": rr})
        flag = f"HIT  rank={rank}" if hit else "miss rank=None"
        print(f"  [{i:2d}/{len(questions)}] {q['type']} {flag}")
    elapsed = time.time() - t0

    n = len(rows)
    hits = sum(1 for r in rows if r["hit"])
    mrr = (sum(r["rr"] for r in rows) / n) if n else 0.0

    print("\n" + "=" * 64)
    print(f"SUMMARY   (k={args.top_k}, {n} questions, {elapsed:.1f}s)")
    print("  retriever : BM25-only (keyword, no semantics)")
    print("=" * 64)
    pct = (100 * hits / n) if n else 0.0
    print(f"  RETRIEVAL   Hit@{args.top_k}: {hits}/{n} = {pct:.1f}%    MRR: {mrr:.3f}")
    print("\n  by type:")
    print(f"    {'type':<18} {'n':>3} {'Hit@k':>7} {'MRR':>7}")
    for t in sorted(set(r["type"] for r in rows)):
        sub = [r for r in rows if r["type"] == t]
        sh = sum(1 for r in sub if r["hit"])
        sm = sum(r["rr"] for r in sub) / len(sub)
        print(f"    {t:<18} {len(sub):>3} {100 * sh / len(sub):>6.0f}% {sm:>7.2f}")

    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = time.strftime("%Y%m%d_%H%M%S")
    out_path = out_dir / f"results_{stamp}_{args.tag}.json"
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(
            {
                "config": {
                    "eval_set": args.eval_set,
                    "bm25_dir": args.bm25_dir,
                    "top_k": args.top_k,
                    "retriever": "bm25_only",
                    "tag": args.tag,
                },
                "summary": {
                    "n": n,
                    "hit_at_k": hits,
                    "hit_rate": (hits / n) if n else 0.0,
                    "mrr": mrr,
                },
                "results": rows,
            },
            f,
            indent=2,
        )
    print(f"\n  saved: {out_path}")


if __name__ == "__main__":
    main()