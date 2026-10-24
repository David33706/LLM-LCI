#!/usr/bin/env python3
"""
step3_orchestrate.py — multi-paper orchestration with TYPE ROUTING.

  TASK1 ("which study contains the LCI for X") -> answered by RETRIEVAL:
        dedup ranking, top-N studies (title + DOI), NO reader call.
  TASK2 ("what assumption was made")           -> full-paper reader (vLLM).

Section trimming and table injection are ON by default (--no-trim / --no-tables to disable).

Run in venv-llmlci, from the repo root, on the same GPU node as the vLLM server:

  python step3_orchestrate.py --qid 8                 # one question
  python step3_orchestrate.py --all --workers 8       # full set (TASK1 routed)
  python step3_orchestrate.py --all --workers 8 --dump-answers  # bare = adopted config (trim + tables ON)
"""

import argparse
import csv
import getpass
import hashlib
import json
import os
import re
import socket
import subprocess
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path

import requests
import openpyxl

_ROOT = os.path.dirname(os.path.abspath(__file__))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

from utils.elsevier import extract_structured_sections, extract_plain_text


# ----------------------------------------------------------------------
# Section trimming (ON by default; --no-trim disables; validated by tools/measure_trim.py)
# ----------------------------------------------------------------------
KEEP_PATTERN = re.compile(
    r"method|material|goal|scope|inventory|\blci\b|\blcia\b|functional unit|"
    r"system boundar|impact assessment|life.?cycle|result|data|case stud|"
    r"scenario|assumption|model(?:ing|ling)?\b",
    re.IGNORECASE,
)
DROP_PATTERN = re.compile(
    r"abstract|introduction|background|literature|related work|discussion|"
    r"conclusion|acknowledg|funding|declaration|author contribution|"
    r"conflict of interest|appendix|supplementary|nomenclature|abbreviation|"
    r"highlights|graphical",
    re.IGNORECASE,
)
MIN_TRIM_TOKENS = 300


def keep_section(section) -> bool:
    stype = (section.get("section_type") or "").lower()
    title = section.get("section_title") or ""
    if stype == "title":
        return True
    if stype == "caption":              # LCI tables live in captions
        return True
    if stype == "abstract":             # removed per Prof. Lam
        return False
    if KEEP_PATTERN.search(title):      # override beats DROP ("Results and discussion")
        return True
    if DROP_PATTERN.search(title):
        return False
    return True                         # unmatched titles kept (safe default)


def build_trimmed_text(xml_text: str):
    """(text, dropped_titles, used_fallback)"""
    sections = extract_structured_sections(xml_text)
    kept_lines, dropped = [], []
    for s in sections:
        t = (s.get("text") or "").strip()
        if not t:
            continue
        if keep_section(s):
            kept_lines.append(f"## {s.get('section_title', 'Section')}\n{t}")
        else:
            dropped.append(s.get("section_title", "?"))
    trimmed = "\n\n".join(kept_lines)
    if count_tokens(trimmed) < MIN_TRIM_TOKENS:   # never produce an unusable feed
        return extract_plain_text(xml_text), dropped, True
    return trimmed, dropped, False


# ----------------------------------------------------------------------
# Table rendering (--tables): inject table CONTENT into the feed.
# extract_structured_sections carries only captions; the cells (the actual
# LCI payload, e.g. the "Primary data" column) never reach the reader
# without this. Rendered verbatim: label + caption + headers + all rows.
# ----------------------------------------------------------------------
def _render_tables(tabs) -> str:
    parts = []
    for t in tabs:
        label = (str(t.get("label") or "").strip()) or "Table"
        cap = str(t.get("caption") or "").strip()
        lines = [f"## {label}: {cap}" if cap else f"## {label}"]
        headers = [str(h) for h in (t.get("headers") or [])]
        if any(h.strip() for h in headers):
            lines.append(" | ".join(headers))
        for row in (t.get("data") or []):
            lines.append(" | ".join("" if c is None else str(c) for c in row))
        parts.append("\n".join(lines))
    return "\n\n".join(parts)


def render_tables_text(xml_path: Path) -> str:
    """All tables of one paper as verbatim text; empty string if none/unparseable."""
    try:
        from tables.xml_tables import extract_tables_from_xml
        tabs = extract_tables_from_xml(xml_path)
    except Exception:
        return ""
    if not tabs:
        return ""
    return _render_tables(tabs)


# ----------------------------------------------------------------------
# Local Qwen tokenizer
# ----------------------------------------------------------------------
_TOK = None
TOKENIZER_SOURCE = None   # path / hub id actually loaded, or "len//4 fallback" (printed in the run header)
def _get_tok():
    global _TOK, TOKENIZER_SOURCE
    if _TOK is None:
        try:
            from transformers import AutoTokenizer
        except Exception:                      # transformers missing -> fall back, don't crash
            AutoTokenizer = None
        work = os.environ.get("LLMLCI_WORK") or f"/work/{getpass.getuser()}"
        for src in (os.path.join(work, "models/Qwen/Qwen2___5-32B-Instruct-AWQ"),  # cluster (offline)
                    "Qwen/Qwen2.5-32B-Instruct-AWQ"):                               # Mac (HF hub)
            if AutoTokenizer is None:
                break
            try:
                _TOK = AutoTokenizer.from_pretrained(src)
                TOKENIZER_SOURCE = src
                break
            except Exception:
                continue
        if _TOK is None:
            _TOK = False
            TOKENIZER_SOURCE = "len//4 fallback"
    return _TOK

def count_tokens(text: str) -> int:
    tok = _get_tok()
    return len(tok(text, add_special_tokens=False)["input_ids"]) if tok else len(text) // 4

def truncate_to_tokens(text: str, max_tokens: int):
    tok = _get_tok()
    if tok:
        ids = tok(text, add_special_tokens=False)["input_ids"]
        if len(ids) <= max_tokens:
            return text, False
        return tok.decode(ids[:max_tokens]), True
    approx = max_tokens * 4
    return (text, False) if len(text) <= approx else (text[:approx], True)


# ----------------------------------------------------------------------
# Reader prompt (TASK2 locator: point + quote, never compute)
# ----------------------------------------------------------------------
SYSTEM_PROMPT = """You are a precise research librarian. You are given the text of ONE scientific life-cycle-assessment (LCA) paper and a question about an assumption or methodological choice made in that paper.

Your job is ONLY to locate and report — never to compute or infer:
- Find where in THIS paper the question is answered.
- Quote the relevant text verbatim: copy it exactly, including any numbers and units as written.
- Do NOT calculate, convert, average, or infer any value. Only report what the text states.
- If the paper does not contain the answer, set "relevant" to false.

Respond with ONLY a JSON object, no other text, in exactly this shape:
{
  "relevant": true,
  "location": "ONLY the section heading or table label under which the quoted text appears (e.g. 'System boundary' or 'Table 2') — copy the heading word-for-word, without any '#' marks, and never write a sentence or description here",
  "assumption": "a short statement of the assumption or choice, in your own words",
  "supporting_quote": "the exact sentence(s) from the paper that contain the answer, copied verbatim"
}
If "relevant" is false, use empty strings for the other three fields."""

_V2_RULES = """A "relevant": true answer must satisfy BOTH rules:

1. SCOPE — this paper's OWN study adopts the assumption for the system the
   question asks about. If the passage only reviews what other studies assume,
   gives general background, or the paper studies a different system or
   process, set "relevant" to false.

2. QUOTE DEPTH — "supporting_quote" is the complete sentence or sentences in
   which the paper STATES the assumption itself, with its specific value,
   number, data source, or named choice. A sentence that merely mentions the
   topic, or points to where the assumption is described, is not sufficient —
   set "relevant" to false instead.

"""
SYSTEM_PROMPT_V2 = SYSTEM_PROMPT.replace("Respond with ONLY a JSON object",
                                         _V2_RULES + "Respond with ONLY a JSON object")
assert SYSTEM_PROMPT_V2 != SYSTEM_PROMPT, "v2 insertion anchor not found"


def doi_to_filename(doi: str) -> str:
    clean = str(doi).replace("/", "_").replace("\\", "_").replace(":", "_")
    return f"{clean}.xml"


def norm_doi(doi: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(doi or "").lower())


EVAL_SET_MISSING = ("ERROR: eval set not found: {}. The project's 50-question eval set is not published; "
                    "pass --eval-set with your own .xlsx (sheet \"Questions\") or .csv with columns "
                    "id,type,question,source_doi (optional: valid_dois, expected_answer, source_location).")


def load_questions(path: str):
    """Eval questions from a .csv or an .xlsx workbook ('Questions' sheet); missing file -> ERROR, exit 1."""
    if not os.path.exists(path):
        sys.exit(EVAL_SET_MISSING.format(path))
    if path.lower().endswith(".csv"):
        # csv: id,type,question,source_doi (+ optional valid_dois, expected_answer, source_location, notes)
        with open(path, newline="", encoding="utf-8-sig") as f:
            rows = [tuple(v if v.strip() else None for v in r) for r in csv.reader(f)]
    else:
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["Questions"] if "Questions" in wb.sheetnames else wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h).strip().lower() if h else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        if r is None or all(v is None for v in r):
            continue
        row = {hdr[i]: (r[i] if i < len(r) else None) for i in range(len(hdr))}
        if (row.get("question") or "") and (row.get("source_doi") or ""):
            if isinstance(row.get("id"), str) and row["id"].strip().isdigit():
                row["id"] = int(row["id"])              # csv ids -> int, as in the workbook
            for k in ("expected_answer", "source_location", "notes"):
                row.setdefault(k, "")                   # absent from the csv -> empty
            row["type"] = str(row.get("type") or "TASK2").strip().upper()
            # optional multi-gold column: "10.1016/a; 10.1016/b" -> normalized set
            raw_valid = str(row.get("valid_dois") or "")
            row["valid_set"] = {norm_doi(d) for d in re.split(r"[;,]", raw_valid) if norm_doi(d)}
            out.append(row)
    return out


def load_titles(wos_csv: str):
    """DOI -> Article Title map (used to answer TASK1 with human-readable titles)."""
    titles = {}
    p = Path(wos_csv)
    if not p.exists():
        return titles
    with open(p, newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            d = norm_doi(row.get("DOI"))
            if d:
                titles[d] = (row.get("Article Title") or "").strip()
    return titles


def dedup_to_papers(chunks, n):
    seen, papers = set(), []
    for ch in chunks:
        raw = (ch.get("metadata", {}) or {}).get("doi")
        k = norm_doi(raw)
        if k and k not in seen:
            seen.add(k)
            papers.append(raw)
            if len(papers) >= n:
                break
    return papers


def derive_location(paper_text: str, quote: str):
    """Find the quote in the feed and return the nearest preceding '## ' heading.
    Model finds the quote; PYTHON resolves the section (same principle as DOIs).
    Whitespace/case-normalized so 'CO 2'-style flattening can't break the match.
    Returns the heading string, or None if the quote isn't verbatim-findable."""
    if not quote or not paper_text:
        return None
    norm_chars, norm_to_orig = [], []
    for i, ch in enumerate(paper_text):
        if not ch.isspace():
            norm_chars.append(ch.lower())
            norm_to_orig.append(i)
    hay = "".join(norm_chars)
    qn = "".join(ch.lower() for ch in quote if not ch.isspace())
    if len(qn) < 20:
        return None
    idx = -1
    for L in (120, 80, 50, 30):
        idx = hay.rfind(qn[:L])   # LAST occurrence: parent sections duplicate child text
        if idx != -1:             # (parent block first, child after) -> rfind lands in the
            break                 # deepest subsection, giving fine-grained headings
    if idx == -1:
        return None
    orig = norm_to_orig[idx]
    h = paper_text.rfind("\n## ", 0, orig)
    if h == -1:
        if paper_text.startswith("## "):
            h = -1  # heading at position 0
        else:
            return None
    start = h + 4  # len("\n## ")
    end = paper_text.find("\n", start)
    heading = paper_text[start:end if end != -1 else None].strip()
    return heading or None


def parse_relevant(value):
    """Strict reading of the reader's "relevant" field. bool("false") is True, so never use bool().
    True/False as is; "true"/"yes" -> True, "false"/"no" -> False (stripped, any case);
    anything else (missing key, 1, "maybe", ...) -> None, which the caller reports as an error."""
    if value is True or value is False:
        return value
    if isinstance(value, str):
        v = value.strip().lower()
        if v in ("true", "yes"):
            return True
        if v in ("false", "no"):
            return False
    return None


def read_one(doi, paper_text, question, url, model, max_tokens, no_think=False, prompt_v2=False):
    user_msg = (
        f"PAPER:\n<<<\n{paper_text}\n>>>\n\n"
        f"QUESTION: {question}\n\n"
        f"Return only the JSON object."
    )
    payload = {
        "model": model,
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT_V2 if prompt_v2 else SYSTEM_PROMPT},
            {"role": "user", "content": user_msg},
        ],
        "temperature": 0,
        "max_tokens": max_tokens,
        "response_format": {"type": "json_object"},
    }
    if no_think:
        payload["chat_template_kwargs"] = {"enable_thinking": False}

    t0 = time.time()
    try:
        resp = requests.post(url, json=payload, timeout=600)
        resp.raise_for_status()
        data = resp.json()
        content = data["choices"][0]["message"]["content"]
        usage = data.get("usage") or {}
    except Exception as e:
        return {"doi": doi, "relevant": None, "error": f"http: {e}", "secs": round(time.time() - t0, 1)}
    if not isinstance(content, str):                   # e.g. None: all output went to reasoning
        return {"doi": doi, "relevant": None, "error": "empty_content",
                "prompt_tokens": usage.get("prompt_tokens"), "secs": round(time.time() - t0, 1)}

    parsed = None
    try:
        parsed = json.loads(content)
    except Exception:
        s, e = content.find("{"), content.rfind("}")
        if s != -1 and e != -1:
            try:
                parsed = json.loads(content[s:e + 1])
            except Exception:
                pass
    if not isinstance(parsed, dict):                   # None, or valid JSON that is not an object
        return {"doi": doi, "relevant": None, "error": "parse_fail",
                "raw": content[:400], "secs": round(time.time() - t0, 1)}
    relevant = parse_relevant(parsed.get("relevant"))
    if relevant is None:                               # missing/odd value: an error, never a YES
        return {"doi": doi, "relevant": None, "error": f"bad_relevant: {parsed.get('relevant')!r}",
                "raw": content[:400], "secs": round(time.time() - t0, 1)}

    loc = str(parsed.get("location", "") or "").strip()
    loc = loc.lstrip("#").strip()                      # model sometimes copies the '## ' feed marker
    loc = " ".join(loc.split())                        # collapse whitespace/newlines
    if len(loc.split()) > 12:                          # heading-length guard: sentences get flagged
        loc = loc + "  [NON-HEADING]"
    loc_source = "model"
    derived = derive_location(paper_text, str(parsed.get("supporting_quote", "") or ""))
    if derived:
        loc, loc_source = derived, "derived"           # computed section beats claimed section
    return {
        "doi": doi,
        "relevant": relevant,
        "location": loc, "location_source": loc_source,
        "assumption": parsed.get("assumption", ""),
        "supporting_quote": parsed.get("supporting_quote", ""),
        "prompt_tokens": usage.get("prompt_tokens"),
        "secs": round(time.time() - t0, 1),
    }


def run_reads(jobs, question, args):
    # ONE argument list for both paths: a new reader parameter added here reaches
    # the sequential (--workers 1) and the ThreadPool path alike.
    def _read(doi, txt):
        return read_one(doi, txt, question, args.url, args.model, args.max_tokens,
                        args.no_think, args.prompt == "v2")

    results = [None] * len(jobs)
    if args.workers <= 1:
        for i, (doi, txt) in enumerate(jobs):
            results[i] = _read(doi, txt)
    else:
        with ThreadPoolExecutor(max_workers=args.workers) as ex:
            futs = {ex.submit(_read, doi, txt): i
                    for i, (doi, txt) in enumerate(jobs)}
            for fut in as_completed(futs):
                results[futs[fut]] = fut.result()
    return results


# ----------------------------------------------------------------------
# TASK1: retrieval answers directly — no reader
# ----------------------------------------------------------------------
def orchestrate_task1(q, pipeline, titles, args, verbose):
    question = q["question"]
    gold_doi = (q.get("source_doi") or "").strip()

    out = pipeline.query(question=question, top_k=args.pool, generate_answer=False)
    chunks = out.get("retrieved_chunks", [])
    papers = dedup_to_papers(chunks, args.num_papers)

    gold_set = {norm_doi(gold_doi)} | q.get("valid_set", set())
    gold_rank = None          # rank of the TAGGED gold (fed-metric stays comparable)
    best_rank = None          # rank of the first ANY-valid paper (multi-gold scoring)
    for i, d in enumerate(papers, 1):
        nd_ = norm_doi(d)
        if gold_rank is None and nd_ == norm_doi(gold_doi):
            gold_rank = i
        if best_rank is None and nd_ in gold_set:
            best_rank = i
    answer = papers[:args.task1_topn]
    gold_fed = gold_rank is not None
    gold_yes = best_rank is not None and best_rank <= args.task1_topn

    if verbose:
        has_gold = bool((q.get("source_doi") or "").strip())
        print("=" * 84)
        print(f"Q{q.get('id')}  |  TASK1 (retrieval-routed, no reader)  |  {question}")
        if has_gold:
            print(f"GOLD: {gold_doi}")
        gr = f"; gold rank = {gold_rank}" if has_gold else ""
        print(f"retrieved {len(chunks)} chunks -> {len(papers)} unique papers; "
              f"answer = top {args.task1_topn}{gr}")
        print("-" * 84)
        for i, d in enumerate(answer, 1):
            star = " <-- GOLD" if norm_doi(d) == norm_doi(gold_doi) else ""
            t = titles.get(norm_doi(d), "(title not in list)")
            print(f"  {i}. {d}{star}\n     {t[:100]}")
        print("-" * 84)
        if has_gold:
            print(f"gold fed: {gold_fed}   gold in answer(top{args.task1_topn}): {gold_yes}")

    return {
        "id": q.get("id"), "type": q["type"], "n_papers": len(papers),
        "n_read": 0, "n_yes": len(answer), "gold_fed": gold_fed,
        "gold_yes": gold_yes, "gold_rank": gold_rank,
        "n_trunc": 0, "n_fallback": 0, "n_loc_derived": 0, "n_loc_total": 0,
        "wall": 0.0, "yes_rows": [],
    }


# ----------------------------------------------------------------------
# TASK2: full-paper reader
# ----------------------------------------------------------------------
def orchestrate_task2(q, pipeline, args, verbose):
    question = q["question"]
    gold_doi = (q.get("source_doi") or "").strip()

    out = pipeline.query(question=question, top_k=args.pool, generate_answer=False)
    chunks = out.get("retrieved_chunks", [])
    papers = dedup_to_papers(chunks, args.num_papers)

    budget = args.ctx - args.max_tokens - 500
    jobs, skipped, ntoks, truncated, trimmed_info, tabinfo = [], [], {}, {}, {}, {}
    for doi in papers:
        p = Path(args.xml_dir) / doi_to_filename(doi)
        if not p.exists():
            skipped.append({"doi": doi, "relevant": None, "error": "xml_missing"})
            continue
        raw = p.read_text(encoding="utf-8", errors="ignore")
        if args.trim:
            txt, dropped, fell_back = build_trimmed_text(raw)
            trimmed_info[doi] = ("fallback" if fell_back else "trim", len(dropped))
        else:
            txt = extract_plain_text(raw)
            trimmed_info[doi] = ("full", 0)
        n_tab = 0
        tbl_block = ""
        if args.tables:
            tbl = render_tables_text(p)
            if tbl:
                n_tab = tbl.count("\n## ") + 1
                tbl_block = "\n\n# TABLES\n\n" + tbl
        tabinfo[doi] = n_tab
        if not txt.strip() and not tbl_block.strip():
            skipped.append({"doi": doi, "relevant": None, "error": "empty_text"})
            continue
        ntoks[doi] = count_tokens(txt + tbl_block)   # TRUE feed size (pre-cut)
        was_trunc = False
        if tbl_block:
            # tables are the LCI payload: reserve their budget FIRST, prose tail pays
            t_tab = count_tokens(tbl_block)
            if t_tab > budget - 4000:                # pathological: tables alone ~fill budget
                tbl_block, _ = truncate_to_tokens(tbl_block, budget - 4000)
                t_tab = count_tokens(tbl_block)
                was_trunc = True
            txt, cut = truncate_to_tokens(txt, budget - t_tab)
            was_trunc = was_trunc or cut
            txt = txt + tbl_block                    # tables appended INTACT
        else:
            txt, was_trunc = truncate_to_tokens(txt, budget)
        truncated[doi] = was_trunc
        jobs.append((doi, txt))

    t0 = time.time()
    read_results = run_reads(jobs, question, args)
    wall = time.time() - t0
    results = read_results + skipped

    # relevant=None never counts as YES; count it by kind so --all runs show it
    reader_errors = {"http": 0, "parse": 0, "other": 0, "xml_missing": 0, "empty_text": 0}
    for r in results:
        if r.get("relevant") is None:
            e = str(r.get("error") or "")
            if e in ("xml_missing", "empty_text"):
                reader_errors[e] += 1
            elif e.startswith("http"):
                reader_errors["http"] += 1
            elif e == "parse_fail":
                reader_errors["parse"] += 1
            else:                                  # bad_relevant, empty_content, ...
                reader_errors["other"] += 1

    yes = [r for r in results if r.get("relevant") is True]
    gold_set = {norm_doi(gold_doi)} | q.get("valid_set", set())
    gold_fed = any(norm_doi(p) == norm_doi(gold_doi) for p in papers)   # tagged only
    gold_yes = any(norm_doi(r["doi"]) in gold_set for r in yes)         # multi-gold
    yes_rows = [{
        "qid": q.get("id"), "type": q["type"], "question": question,
        "doi": r["doi"],
        "is_tagged_gold": "YES" if norm_doi(r["doi"]) == norm_doi(gold_doi) else "",
        "location": r.get("location", ""), "location_source": r.get("location_source", ""),
        "assumption": r.get("assumption", ""),
        "supporting_quote": r.get("supporting_quote", ""),
    } for r in yes]

    if verbose:
        n_tr = sum(1 for v in truncated.values() if v)
        print("=" * 84)
        print(f"Q{q.get('id')}  |  TASK2 (reader)  |  {question}")
        if (q.get("source_doi") or "").strip():
            print(f"GOLD: {gold_doi}   ({q.get('source_location')})")
        print(f"retrieved {len(chunks)} chunks -> {len(papers)} unique papers; "
              f"{len(jobs)} read ({n_tr} truncated), {len(skipped)} skipped; wall {wall:.1f}s "
              f"(workers={args.workers}, trim={'on' if args.trim else 'off'})")
        print("-" * 84)
        for r in results:
            star = " <-- GOLD" if norm_doi(r["doi"]) == norm_doi(gold_doi) else ""
            nt = ntoks.get(r["doi"])
            tag = f"{nt:>6} tok" if nt is not None else "     ? tok"
            mode, ndropped = trimmed_info.get(r["doi"], ("?", 0))
            tinfo = {"trim": f" trim(-{ndropped} sect)", "fallback": " FALLBACK", "full": ""}.get(mode, "")
            if args.tables:
                tinfo += f" +{tabinfo.get(r['doi'], 0)}tab"
            tr = " [TRUNC]" if truncated.get(r["doi"]) else ""
            if r.get("error"):
                print(f"  [{tag}]{tinfo}{tr} {r['doi']}{star}   ERR: {str(r['error'])[:42]}")
                continue
            mark = "YES" if r["relevant"] else " . "
            print(f"  [{tag}]{tinfo}{tr} [{mark}] {r['doi']}{star}")
            if r["relevant"]:
                print(f"        loc : {r['location']}")
                print(f"        quote: {str(r['supporting_quote'])[:160]}")
        print("-" * 84)
        print(f"ANSWER: {len(yes)} paper(s) contain this — " +
              (", ".join(r["doi"] for r in yes) if yes else "none"))
        if (q.get("source_doi") or "").strip():
            print(f"gold fed: {gold_fed}   gold in answer: {gold_yes}")

    return {
        "id": q.get("id"), "type": q["type"], "n_papers": len(papers),
        "n_read": len(jobs), "n_yes": len(yes),
        "gold_fed": gold_fed, "gold_yes": gold_yes, "gold_rank": None,
        "n_trunc": sum(1 for v in truncated.values() if v),
        "n_fallback": sum(1 for m, _ in trimmed_info.values() if m == "fallback"),
        "n_loc_derived": sum(1 for r in yes if r.get("location_source") == "derived"),
        "n_loc_total": len(yes),
        "wall": wall, "yes_rows": yes_rows,
        "reader_errors": reader_errors,
    }


def orchestrate(q, pipeline, titles, args, verbose):
    if q["type"].startswith("TASK1"):
        return orchestrate_task1(q, pipeline, titles, args, verbose)
    return orchestrate_task2(q, pipeline, args, verbose)


def build_pipeline(args):
    """HYBRID base at the validated config (weighted, alpha=0.4). No cross-encoder:
    whole-paper feed only needs the paper SET, which equals the hybrid ranking."""
    from hybrid import build_hybrid_retriever
    return build_hybrid_retriever(
        chroma_dir=args.chroma_dir,
        chroma_collection=args.collection,
        embedding_model=args.embedding_model,
        bm25_dir=args.bm25_dir,
        llm_model=None,
        rrf_k=60,
        candidate_k=args.candidate_k,
        fusion="weighted",              # pinned: validated config
        alpha=args.alpha,               # default 0.4 (validated)
    )

def build_parser():
    ap = argparse.ArgumentParser()
    ap.add_argument("--qid", type=int, help="run a single question by id")
    ap.add_argument("--all", action="store_true", help="run the whole eval set")
    ap.add_argument("--eval-set", default="eval/LCA_RAG_eval_set_v3_50.xlsx",
                    help="eval questions: .xlsx or .csv (default: the private v3_50 workbook, not published)")
    ap.add_argument("--xml-dir", default="elsevier_xml")
    ap.add_argument("--wos", default="paper_list2_2025.csv", help="DOI->title map for TASK1 answers")
    # routing / trimming
    ap.add_argument("--task1-topn", type=int, default=3,
                    help="TASK1 answer = top-N studies; gold_yes = gold within N")
    ap.add_argument("--dump-answers", action="store_true",
                    help="with --all: write every TASK2 yes-answer (doi/location/quote) to "
                         "eval_runs/answers_<ts>.xlsx for human validation (valid_dois workflow)")
    ap.add_argument("--tables", action=argparse.BooleanOptionalAction, default=True,
                    help="inject rendered table content into the feed (--no-tables to disable)")
    ap.add_argument("--trim", action=argparse.BooleanOptionalAction, default=True,
                    help="section trimming (--no-trim to disable)")
    # reader (vLLM)
    ap.add_argument("--url", default="http://localhost:3000/v1/chat/completions")
    ap.add_argument("--model", default="qwen3.5-35b-a3b-fp8")
    ap.add_argument("--max-tokens", type=int, default=1500)
    ap.add_argument("--workers", type=int, default=None,
                    help="reader calls in flight (default: = --num-papers; 1 = sequential)")
    # orchestration
    ap.add_argument("--num-papers", type=int, default=8)
    ap.add_argument("--pool", type=int, default=40)
    # retriever (pinned validated config)
    ap.add_argument("--chroma-dir", default="chroma_db")
    ap.add_argument("--collection", default="papers")
    ap.add_argument("--embedding-model", default="sentence-transformers/all-MiniLM-L6-v2")
    ap.add_argument("--bm25-dir", default=None)
    ap.add_argument("--alpha", type=float, default=0.4)
    ap.add_argument("--candidate-k", type=int, default=100)

    ap.add_argument("--ctx", type=int, default=131072,
                    help="server max-model-len; feed budget = ctx - max-tokens - 500")
    ap.add_argument("--no-think", action="store_true",
                    help="force-disable thinking (auto-set for qwen3* models)")
    ap.add_argument("--think", action="store_true",
                    help="force-enable thinking even for qwen3* models")
    ap.add_argument("--prompt", choices=["v1", "v2"], default="v1",
                    help="reader prompt variant; v2 adds scope + quote-depth rules")

    return ap

def resolve_args(args):
    if args.think:
        args.no_think = False
    elif not args.no_think:
        args.no_think = args.model.lower().startswith("qwen3")
    if args.workers is None:
        args.workers = args.num_papers


def write_manifest(path, args, n_store, titles, out_path, t_start):
    """Run manifest next to the answers dump: what ran, on which data, where, when."""
    def _safe(v):
        return v if isinstance(v, (str, int, float, bool, type(None))) else str(v)
    with open(args.eval_set, "rb") as f:
        eval_sha = hashlib.sha256(f.read()).hexdigest()
    try:
        gp = subprocess.run(["git", "rev-parse", "HEAD"], cwd=_ROOT,
                            capture_output=True, text=True, timeout=10)
        commit = gp.stdout.strip() if gp.returncode == 0 and gp.stdout.strip() else None
    except Exception:
        commit = None                          # cluster checkout is rsynced, no git
    manifest = {
        "argv": sys.argv,
        "args": {k: _safe(v) for k, v in vars(args).items()},
        "vector_store_count": n_store,
        "tokenizer_source": TOKENIZER_SOURCE,
        "titles_loaded": len(titles), "titles_path": args.wos,
        "eval_set": args.eval_set, "eval_set_sha256": eval_sha,
        "answers_dump": out_path,
        "git_commit": commit,
        "hostname": socket.gethostname(),
        "slurm_job_id": os.environ.get("SLURM_JOB_ID"),
        "start_time": t_start.isoformat(timespec="seconds"),
        "end_time": datetime.now().isoformat(timespec="seconds"),
    }
    with open(path, "w", encoding="utf-8") as f:
        json.dump(manifest, f, indent=2)
        f.write("\n")


def main():
    t_start = datetime.now()
    ap = build_parser()
    args = ap.parse_args()
    resolve_args(args)

    if not args.all and args.qid is None:
        ap.error("pass --qid N or --all")

    questions = load_questions(args.eval_set)
    by_id = {str(q.get("id")): q for q in questions}
    titles = load_titles(args.wos)

    print(f"building retriever (hybrid, fusion=weighted, alpha={args.alpha}, "
          f"candidate_k={args.candidate_k}) | trim={'ON' if args.trim else 'OFF'} | "
          f"tables={'ON' if args.tables else 'OFF'} | "
          f"TASK1 -> retrieval top-{args.task1_topn}")
    pipeline = build_pipeline(args)
    try:
        n_store = pipeline.vector_store.count()
    except Exception as e:
        sys.exit(f"ERROR: vector store is empty or unreadable ({type(e).__name__}: {e}). "
                 f"Run from the repo root and check chroma_db/ (RUNBOOK §4).")
    print(f"  vector store: {n_store} chunks")
    if not n_store:
        sys.exit(f"ERROR: vector store is empty or unreadable (0 chunks in '{args.chroma_dir}', "
                 f"collection '{args.collection}'). Run from the repo root and check chroma_db/ (RUNBOOK §4).")
    _get_tok()
    print(f"  tokenizer: {TOKENIZER_SOURCE}")
    if TOKENIZER_SOURCE == "len//4 fallback":
        print("  WARNING: token counts use len//4 fallback (no Qwen tokenizer found); "
              "truncation and budget are approximate")
    print(f"  titles: {len(titles)} loaded from {args.wos}")
    if not titles:
        print(f"  WARNING: 0 titles loaded from {args.wos}; TASK1 answers will show (title not in list)")

    if args.all:
        rows, t0 = [], time.time()
        for i, q in enumerate(questions, 1):
            print(f"\n[{i}/{len(questions)}] Q{q.get('id')} ({q['type']}) ...", flush=True)
            rows.append(orchestrate(q, pipeline, titles, args, verbose=False))
        total = time.time() - t0
        n = len(rows)
        t1 = [r for r in rows if r["type"].startswith("TASK1")]
        t2 = [r for r in rows if not r["type"].startswith("TASK1")]
        gf = sum(r["gold_fed"] for r in rows)
        gy = sum(r["gold_yes"] for r in rows)
        print("\n" + "=" * 70)
        print(f"SUMMARY  ({n} questions, {total:.0f}s total, workers={args.workers}, "
              f"trim={'ON' if args.trim else 'OFF'}, tables={'ON' if args.tables else 'OFF'})")
        n_mg = sum(1 for q_ in questions if q_.get("valid_set"))
        mg_note = f"   [multi-gold active on {n_mg} question(s)]" if n_mg else ""
        print(f"  COMPOSITE  gold fed: {gf}/{n} = {gf/n:.0%}   gold in answer: {gy}/{n} = {gy/n:.0%}{mg_note}")
        if t1:
            g1 = sum(r["gold_yes"] for r in t1)
            g1f = sum(r["gold_fed"] for r in t1)
            r1 = [r["gold_rank"] for r in t1 if r["gold_rank"] == 1]
            print(f"  TASK1 (retrieval, top-{args.task1_topn}): {g1}/{len(t1)}   "
                  f"(fed {g1f}/{len(t1)}; rank-1: {len(r1)})")
        if t2:
            g2 = sum(r["gold_yes"] for r in t2)
            g2f = sum(r["gold_fed"] for r in t2)
            print(f"  TASK2 (reader)              : {g2}/{len(t2)}   (fed {g2f}/{len(t2)})")
        print(f"  avg papers read/question (TASK2): "
              f"{(sum(r['n_read'] for r in t2)/len(t2)) if t2 else 0:.1f}")
        ld = sum(r.get("n_loc_derived", 0) for r in rows)
        lt = sum(r.get("n_loc_total", 0) for r in rows)
        if lt:
            print(f"  locations python-verified       : {ld}/{lt} yes-answers")
        print(f"  papers truncated (too long)     : {sum(r['n_trunc'] for r in rows)}")
        if args.trim:
            print(f"  trim fallbacks                  : {sum(r['n_fallback'] for r in rows)}")
        print(f"  avg wall/question (TASK2)       : "
              f"{(sum(r['wall'] for r in t2)/len(t2)) if t2 else 0:.1f}s")
        ec = {k: sum(r.get("reader_errors", {}).get(k, 0) for r in rows)
              for k in ("http", "parse", "other", "xml_missing", "empty_text")}
        n_err = ec["http"] + ec["parse"] + ec["other"]
        print(f"  reader errors: {n_err} (http {ec['http']}, parse {ec['parse']}, other {ec['other']}); "
              f"skipped papers: xml_missing {ec['xml_missing']}, empty_text {ec['empty_text']}")
        if n_err:
            bad = [str(r["id"]) for r in rows
                   if any(r.get("reader_errors", {}).get(k, 0) for k in ("http", "parse", "other"))]
            print(f"  WARNING: {n_err} reader call(s) failed and counted as not-YES; "
                  f"re-run with --qid N to see the ERR lines (qids: {', '.join(bad)})")
        print("=" * 70)
        if args.dump_answers:
            os.makedirs("eval_runs", exist_ok=True)
            ts = datetime.now().strftime('%Y%m%d_%H%M%S')
            out_path = os.path.join("eval_runs", f"answers_{ts}.xlsx")
            wb_out = openpyxl.Workbook()
            ws_out = wb_out.active
            ws_out.title = "yes_answers"
            cols = ["qid", "type", "question", "doi", "is_tagged_gold",
                    "location", "location_source", "assumption", "supporting_quote", "valid?"]
            ws_out.append(cols)
            for r in rows:
                for yr in r.get("yes_rows", []):
                    ws_out.append([yr.get(c, "") for c in cols[:-1]] + [""])
            widths = [6, 8, 50, 30, 13, 28, 12, 45, 70, 8]
            from openpyxl.utils import get_column_letter
            for i, w in enumerate(widths, 1):
                ws_out.column_dimensions[get_column_letter(i)].width = w
            ws_out.freeze_panes = "A2"
            wb_out.save(out_path)
            n_rows = sum(len(r.get("yes_rows", [])) for r in rows)
            print(f"\n  answers dump: {out_path}  ({n_rows} yes-rows; fill the 'valid?' column)")
            man_path = os.path.join("eval_runs", f"manifest_{ts}.json")
            try:
                write_manifest(man_path, args, n_store, titles, out_path, t_start)
                print(f"  manifest: {man_path}")
            except Exception as e:                 # the manifest must never break a run
                print(f"  WARNING: manifest not written ({type(e).__name__}: {e})")
        print("  per question (id/type: fed/read/yes  goldFed goldYes  rank  trunc):")
        for r in rows:
            rk = r["gold_rank"] if r["gold_rank"] is not None else "-"
            print(f"    {str(r['id']):>3}/{r['type'][:5]:<5}: {r['n_papers']}/{r['n_read']}/{r['n_yes']}"
                  f"   {'Y' if r['gold_fed'] else '.'}      {'Y' if r['gold_yes'] else '.'}"
                  f"      {rk:>3}   {r['n_trunc'] or '.'}")
    else:
        q = by_id.get(str(args.qid))
        if not q:
            ap.error(f"qid {args.qid} not in {args.eval_set}")
        orchestrate(q, pipeline, titles, args, verbose=True)


if __name__ == "__main__":
    main()