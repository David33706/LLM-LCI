#!/usr/bin/env python3
"""
index_tables.py (v2) - add table records to the 'papers' collection, now with the
paper TITLE prepended to each table-document.

WHY THE TITLE:
TASK1 queries describe the STUDY ("Legionella disinfection in hospitals") plus the
data type ("annual consumption inventory"). The bare table-document has the data-type
words (from caption/headers) but NOT the study-context words, so it ranks below the
prose that does. Prepending the paper title injects those terms. This is the ONLY change
vs v1 (one variable), so any move in right-table Hit@5 is attributable to it.

  document = [DOI: <doi>] <TITLE>. <label>. <caption> Columns: ... Entries: ...
             (numbers mostly, but not completely, excluded: is_number() drops plain
              numeric cells but misses forms like '5.3 × 10⁻⁴' or '1.2 ± 0.3')
  embedding = YOUR MiniLM over that text (explicit -> same space as prose)
  metadata  = {doi, section_title=<label>, section_type="table",
               type="table", table_number=<label>, caption, markdown=<full grid>}
  id        = "<doi>::table::<label>"  (deterministic -> upsert overwrites in place)

SAFE + REVERSIBLE:
  - upsert by stable id -> re-running OVERWRITES the same records (no dupes);
    current corpus: 9,994 table records from 1,824 papers
  - to REVERT to the no-title version:  add --no-title and re-run --all
  - to remove tables entirely: coll.delete(where={"type":"table"}) then rebuild BM25

USAGE (from the repo root)
  python ingest/index_tables.py --all --wos paper_list2_2025.csv  # full corpus, title-enriched
  python ingest/index_tables.py --all --no-title                  # revert to v1 behavior
Then: rebuild BM25, then re-run eval/table_retrieval_eval.py.

GUARDS (fail loudly, exit 1)
  - titles on (no --no-title) but the --wos file is missing or loads 0 titles
    (the documented '0 titles loaded' null run);
  - --xlsx-dir missing, or nothing to index;
  - after upserting, the count delta differs from the number of built records
    whose ids were not already in the collection.
"""
import argparse, csv,  re
import openpyxl
from xml.etree import ElementTree as ET

import os, sys
_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

csv.field_size_limit(10_000_000)
DEFAULT_DOI = "10.1016/j.afres.2025.100716"


def doi_value_from_stem(stem):
    return stem.replace("_", "/", 1)

def norm_key(s):
    return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")

def norm_doi(s):
    return re.sub(r"[^a-z0-9]", "", (s or "").lower())

def find_file(directory, doi_value, ext):
    if not os.path.isdir(directory):
        return None
    target = norm_key(doi_value)
    for f in sorted(os.listdir(directory)):
        stem, e = os.path.splitext(f)
        if e.lower() == ext and norm_key(stem) == target:
            return os.path.join(directory, f)
    return None

def is_number(c):
    return bool(re.fullmatch(r"[\d.,\s×x^+\-eE()%]*", c or ""))

def load_titles(wos_path):
    """ {normalized_doi: Article Title} from the DOI->title list; {} if file missing. """
    titles = {}
    if not wos_path or not os.path.exists(wos_path):
        return titles
    with open(wos_path, newline="", encoding="utf-8-sig", errors="replace") as f:
        for row in csv.DictReader(f):
            k = norm_doi(row.get("DOI"))
            if k:
                titles[k] = (row.get("Article Title") or "").strip()
    return titles

def load_grids(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    return {s: [["" if c is None else str(c).strip() for c in r]
                for r in wb[s].iter_rows(values_only=True)] for s in wb.sheetnames}

def captions_from_xml(xml_path):
    caps = {}
    if not xml_path or not os.path.exists(xml_path):
        return caps
    try:
        root = ET.parse(xml_path).getroot()
    except Exception:
        return caps
    for el in root.iter():
        if el.tag.split('}')[-1] != "table":
            continue
        label = cap = None
        for ch in el.iter():
            t = ch.tag.split('}')[-1]
            if t == "label" and label is None:
                label = " ".join("".join(ch.itertext()).split())
            elif t == "caption" and cap is None:
                cap = " ".join("".join(ch.itertext()).split())
        if label:
            caps[label.strip()] = (cap or "").strip()
    return caps

def build_document(doi, title, label, caption, rows):
    headers = [c for c in rows[0] if c] if rows else []
    entities = list(dict.fromkeys(c for r in rows[1:] for c in r if c and not is_number(c) and len(c) < 40))
    parts = [f"[DOI: {doi}]"]
    if title:
        parts.append(title + ".")
    parts.append(f"{label}.")
    if caption:
        parts.append(caption)
    parts += ["Columns: " + ", ".join(headers) + ".", "Entries: " + ", ".join(entities) + "."]
    return re.sub(r"\s+", " ", " ".join(parts)).strip()

def grid_markdown(rows):
    return "\n".join("| " + " | ".join(r) + " |" for r in rows)

def records_for_paper(doi_value, title, xlsx_dir, xml_dir):
    xlsx = find_file(xlsx_dir, doi_value, ".xlsx")
    if not xlsx:
        return []
    grids = load_grids(xlsx)
    caps = captions_from_xml(find_file(xml_dir, doi_value, ".xml"))
    recs = []
    for label, rows in grids.items():
        caption = caps.get(label, "")
        doc = build_document(doi_value, title, label, caption, rows)
        meta = {"doi": doi_value, "section_title": label, "section_type": "table",
                "type": "table", "table_number": label, "caption": caption,
                "markdown": grid_markdown(rows)}
        recs.append((f"{doi_value}::table::{label}", doc, meta))
    return recs

def main():
    ap = argparse.ArgumentParser(description="Index table records (title-enriched) into papers.")
    ap.add_argument("--all", action="store_true")
    ap.add_argument("--doi", default=DEFAULT_DOI)
    ap.add_argument("--xlsx-dir", default="tables_xml")
    ap.add_argument("--xml-dir", default="elsevier_xml")
    ap.add_argument("--wos", default="paper_list2_2025.csv", help="DOI->title list (corpus DOIs, titles from Crossref)")
    ap.add_argument("--no-title", action="store_true", help="Disable title prefix (reverts to v1)")
    ap.add_argument("--chroma-dir", default="chroma_db")
    ap.add_argument("--collection", default="papers")
    ap.add_argument("--embedding-model", default="sentence-transformers/all-MiniLM-L6-v2")
    ap.add_argument("--batch", type=int, default=256)
    args = ap.parse_args()

    titles = {} if args.no_title else load_titles(args.wos)
    print(f"title prefix: {'OFF (--no-title)' if args.no_title else f'ON ({len(titles)} titles loaded from {args.wos})'}")
    if not args.no_title and not titles:
        why = "file not found" if not os.path.exists(args.wos) else "0 titles loaded"
        print(f"ERROR: --wos {args.wos}: {why}. Table records would be indexed without "
              f"titles (the '0 titles loaded' null run). Pass --wos paper_list2_2025.csv "
              f"from the repo root, or --no-title on purpose.")
        sys.exit(1)

    if args.all:
        if not os.path.isdir(args.xlsx_dir):
            print(f"ERROR: --xlsx-dir not found: {args.xlsx_dir}"); sys.exit(1)
        dois = [doi_value_from_stem(os.path.splitext(f)[0])
                for f in sorted(os.listdir(args.xlsx_dir)) if f.lower().endswith(".xlsx")]
        print(f"full build: {len(dois)} papers with tables")
    else:
        dois = [args.doi]
        print(f"single paper: {args.doi}")

    ids, docs, metas = [], [], []
    n_papers = n_titled = 0
    for d in dois:
        title = titles.get(norm_doi(d), "")
        recs = records_for_paper(d, title, args.xlsx_dir, args.xml_dir)
        if recs:
            n_papers += 1
            if title:
                n_titled += 1
        for rid, doc, meta in recs:
            ids.append(rid); docs.append(doc); metas.append(meta)
    print(f"built {len(ids)} table records from {n_papers} papers "
          f"({n_titled} had a matched title, {n_papers - n_titled} did not)")
    if not ids:
        print("nothing to index. Check --doi / --xlsx-dir.")
        print("ERROR: 0 table records built"); sys.exit(1)

    from rag import build_default_pipeline
    pipe = build_default_pipeline(chroma_dir=args.chroma_dir, chroma_collection=args.collection,
                                  embedding_model=args.embedding_model, llm_model=None)
    pipe.vector_store._ensure_collection()
    coll = pipe.vector_store._collection
    before = coll.count()
    print(f"collection '{args.collection}' count before: {before}")
    preexisting = 0                     # ids already present -> upsert overwrites, adds 0
    for i in range(0, len(ids), 1000):  # paged: large get() hits "too many SQL variables"
        preexisting += len(coll.get(ids=ids[i:i+1000], include=[]).get("ids", []))
    expected_new = len(ids) - preexisting
    print(f"example document:\n  {docs[0][:160]}...")

    embedder = pipe.embedder
    for i in range(0, len(ids), args.batch):
        bd = docs[i:i+args.batch]
        embs = embedder.embed_texts(bd)
        coll.upsert(ids=ids[i:i+args.batch], documents=bd,
                    metadatas=metas[i:i+args.batch], embeddings=embs.tolist())
        print(f"  upserted {min(i+args.batch, len(ids))}/{len(ids)}")

    after = coll.count()
    print(f"collection '{args.collection}' count after : {after}   (+{after - before}; 0 expected if overwriting)")
    print(f"records built: {len(ids)} | already present (overwritten): {preexisting} | "
          f"expected new: {expected_new}")
    if after - before != expected_new:
        print(f"ERROR: count grew by {after - before} but {expected_new} new table records "
              f"were upserted ({len(ids)} built, {preexisting} already present). "
              f"Duplicate/colliding ids or a concurrent writer; inspect the collection.")
        sys.exit(1)
    print("\nNEXT:")
    print("  python ingest/build_bm25_index.py --collection papers --out bm25_index/papers")
    print("  eval/table_retrieval_eval.py needs the full private eval workbook (v2 answers),")
    print("  which is not published. With that workbook:")
    print("  PYTHONPATH=. python eval/table_retrieval_eval.py --eval-set FULL_WORKBOOK.xlsx \\")
    print("      --retriever hybrid --fusion weighted --alpha 0.4 --bm25-dir bm25_index/papers")

if __name__ == "__main__":
    main()